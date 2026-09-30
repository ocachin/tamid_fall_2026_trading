"""ETF arbitrage monitor in Excel (xlwings).

Python is only the data feed: it writes raw quotes and historical stats to the Feed sheet.
Every number on the Dashboard is an Excel formula on Feed, Holdings and Config.

    .venv/bin/python etf_arb/tracker.py                # stream live
    .venv/bin/python etf_arb/tracker.py --once         # one refresh
    .venv/bin/python etf_arb/tracker.py holdings       # re-download all holdings
    .venv/bin/python etf_arb/tracker.py setup --force  # rebuild the workbook
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import os
import re
import sys
import time
import traceback
from pathlib import Path

os.environ.setdefault("MPLBACKEND", "Agg")

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import xlwings as xw
import yfinance as yf
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from arb import ET, Model, Settings, build_model, history_frame, session_left  # noqa: E402
from data import Alpaca, Nasdaq, Yahoo, fetch_holdings, yahoo_symbol  # noqa: E402

BOOK = HERE / "ETF_Arb_Tracker.xlsx"
LOGS = HERE / "logs"
CACHE = HERE / "cache"
SHEETS = ["Dashboard", "Universe", "Holdings", "Charts", "Chart Data", "Config", "Feed"]
MAX_ROW = 600

plt.rcParams.update({"font.family": "Arial", "font.size": 9, "axes.spines.top": False,
                     "axes.spines.right": False, "axes.grid": False, "axes.titlesize": 10,
                     "axes.titleweight": "bold", "legend.frameon": False})

# ------------------------------------------------------------------ inputs
UNIVERSE_COLS = ["ETF", "Class", "Model", "Underlying", "Leverage", "Peers", "Holdings Source", "Active"]
DEFAULT_UNIVERSE = [
    ["SPY", "EQUITY", "BASKET", "", "", "IVV,VOO", "SPDR", "Y"],
    ["QQQ", "EQUITY", "BASKET", "", "", "QQQM", "INVESCO", "Y"],
    ["XLK", "EQUITY", "BASKET", "", "", "VGT", "SPDR", "Y"],
    ["XLF", "EQUITY", "BASKET", "", "", "VFH", "SPDR", "Y"],
    ["XLE", "EQUITY", "BASKET", "", "", "VDE", "SPDR", "Y"],
    ["XLV", "EQUITY", "BASKET", "", "", "VHT", "SPDR", "Y"],
    ["IWM", "EQUITY", "PROXY", "VTWO", "", "VTWO", "", "Y"],
    ["AAPU", "EQUITY", "PROXY", "AAPL", 2, "", "", "Y"],
    ["TLT", "FI", "PROXY", "VGLT", "", "VGLT,SPTL", "", "Y"],
    ["IEF", "FI", "PROXY", "VGIT", "", "VGIT,SCHR", "", "Y"],
    ["SHY", "FI", "PROXY", "VGSH", "", "VGSH,SCHO", "", "Y"],
    ["LQD", "FI", "PROXY", "IEF", "", "VCIT,IGIB", "", "Y"],
    ["VCIT", "FI", "PROXY", "IEF", "", "IGIB", "", "Y"],
    ["HYG", "FI", "PROXY", "JNK", "", "JNK,USHY", "", "Y"],
    ["AGG", "FI", "PROXY", "BND", "", "BND,SCHZ", "", "Y"],
    ["TIP", "FI", "PROXY", "SCHP", "", "SCHP,SPIP", "", "Y"],
    ["MUB", "FI", "PROXY", "VTEB", "", "VTEB", "", "Y"],
]

# label, Settings field, default, Excel name (for the settings the Dashboard formulas use)
CONFIG = [
    ("Live", "live", "Y", None),
    ("Refresh (sec)", "refresh_seconds", 20, None),
    ("Equity cost (bps)", "equity_cost_bps", 3, "CostEq"),
    ("FI cost (bps)", "fi_cost_bps", 8, "CostFI"),
    ("Equity min edge (bps)", "equity_min_edge_bps", 2, "EdgeEq"),
    ("FI min edge (bps)", "fi_min_edge_bps", 5, "EdgeFI"),
    ("Anchor (CLOSE/NAV)", "anchor", "CLOSE", "Anchor"),
    ("Stale after (min)", "stale_minutes", 20, "StaleMin"),
    ("Ex-div flag (bps)", "exdiv_flag_bps", 30, "ExDivFlag"),
    ("Holdings shown", "holdings_shown", 3, None),
    ("Max basket names", "max_basket_names", 600, None),
    ("History period", "history_period", "1y", None),
    ("Beta lookback (days)", "beta_lookback", 120, None),
    ("Resid lookback (days)", "z_lookback", 60, None),
    ("Vol lookback (days)", "vol_lookback", 60, None),
    ("Chart refresh (sec)", "chart_seconds", 120, None),
    ("Log intraday", "log_live", "Y", None),
]

# ------------------------------------------------------------------ Feed + Dashboard layout
FEED_COLS = ["Symbol", "Last", "Bid", "Ask", "Bid Sz", "Ask Sz", "Prev Close", "Volume", "Avg Vol",
             "NAV", "Yield", "Sigma", "Time", "Beta", "Resid Mean", "Resid SD"]
F = {name: get_column_letter(i + 1) for i, name in enumerate(FEED_COLS)}      # Feed column letters

PX, QTY = "#,##0.00;-#,##0.00;;@", "#,##0;-#,##0;;@"
DASH = [  # header, width, number format
    ("Ticker", 9, "General"), ("Signal", 21, "General"), ("β / Wt", 8, "0.0%;-0.0%;;@"),
    ("Last", 9, PX), ("Bid", 9, PX), ("Ask", 9, PX), ("Bid Sz", 8, QTY), ("Ask Sz", 8, QTY),
    ("Spread bps", 8, "0.0"), ("Chg %", 8, "0.00%"), ("Volume", 12, QTY), ("Avg Vol", 12, QTY),
    ("Rel Vol", 7, '0.00"x"'), ("σ 1D", 7, "0.00%;-0.00%;;@"), ("Move σ", 7, "0.00"),
    ("P>1σ", 7, "0%"), ("P>2σ", 7, "0%"), ("P>3σ", 7, "0%"),
    ("NAV T-1", 9, PX), ("Prem T-1 bps", 9, "0.0"), ("Und Move", 8, "0.00%"),
    ("Fair Value", 10, "#,##0.000"), ("vs FV bps", 8, "0.0"), ("Sell Edge", 8, "0.0"),
    ("Buy Edge", 8, "0.0"), ("Resid Z", 7, "0.00"), ("Yield", 7, "0.00%;-0.00%;;@"),
    ("Time", 9, "hh:mm:ss;;;@"),
]
C = {h: get_column_letter(i + 1) for i, (h, _, _) in enumerate(DASH)}          # Dashboard column letters
HIDDEN = {"Key": "AD", "Role": "AE", "Class": "AF", "Row": "AG"}
HDR, FIRST = 2, 3                                                        # header row, first data row

FONT = "Arial"
THIN = Side(style="thin", color="000000")
HAIR = Side(style="thin", color="BFBFBF")
INPUT_BLUE = "0000FF"


# ------------------------------------------------------------------ workbook build (openpyxl)
def build_workbook() -> None:
    """Static layout, formats, dropdowns and conditional formatting. Formulas are written live."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = {n: wb.create_sheet(n) for n in SHEETS}
    for sh in ws.values():
        sh.sheet_view.showGridLines = False
    base, bold = Font(name=FONT, size=10), Font(name=FONT, size=10, bold=True)
    blue = Font(name=FONT, size=10, color=INPUT_BLUE)

    def header(sh, row: int, labels: list[str], widths: list[float], right_from: int = 99):
        for j, lab in enumerate(labels, start=1):
            c = sh.cell(row, j, lab or None)
            c.font, c.border = bold, Border(bottom=THIN)
            c.alignment = Alignment(horizontal="right" if j >= right_from else "left",
                                    vertical="bottom", wrap_text=True)
            sh.column_dimensions[c.column_letter].width = widths[j - 1]

    # Dashboard
    d = ws["Dashboard"]
    d["A1"].font = Font(name=FONT, size=8, color="808080")
    header(d, HDR, [h for h, _, _ in DASH], [w for _, w, _ in DASH], right_from=3)
    d.row_dimensions[HDR].height = 27
    for j, (_, _, fmt) in enumerate(DASH, start=1):
        for r in range(FIRST, MAX_ROW + 1):
            c = d.cell(r, j)
            c.number_format, c.font = fmt, base
    for col in HIDDEN.values():
        d.column_dimensions[col].hidden = True
    d.freeze_panes = f"B{FIRST}"
    role = f'${HIDDEN["Role"]}{FIRST}'
    d.conditional_formatting.add(f"A{FIRST}:{C['Time']}{MAX_ROW}", FormulaRule(
        formula=[f'{role}="ETF"'], font=Font(bold=True), border=Border(top=HAIR)))
    d.conditional_formatting.add(f"C{FIRST}:C{MAX_ROW}", Rule(
        type="expression", formula=[f'{role}="ETF"'],
        dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=200, formatCode='0.00"x"'))))
    for word, font, fill in (("SELL", "9C0006", "F4CCCC"), ("BUY", "006100", "D9EAD3")):
        d.conditional_formatting.add(f"B{FIRST}:B{MAX_ROW}", FormulaRule(
            formula=[f'LEFT(B{FIRST},{len(word)})="{word}"'], font=Font(color=font),
            fill=PatternFill("solid", start_color=fill, end_color=fill)))
    d.conditional_formatting.add(f"{C['P>1σ']}{FIRST}:{C['P>3σ']}{MAX_ROW}", ColorScaleRule(
        start_type="num", start_value=0, start_color="FFFFFF",
        mid_type="num", mid_value=0.25, mid_color="FFD966",
        end_type="num", end_value=1, end_color="E06666"))
    d.conditional_formatting.add(f"{C['Move σ']}{FIRST}:{C['Move σ']}{MAX_ROW}", ColorScaleRule(
        start_type="num", start_value=-3, start_color="E06666",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=3, end_color="6AA84F"))
    for col in ("Sell Edge", "Buy Edge"):
        d.conditional_formatting.add(f"{C[col]}{FIRST}:{C[col]}{MAX_ROW}", CellIsRule(
            operator="greaterThan", formula=["0"], font=Font(color="006100", bold=True)))

    # Universe
    u = ws["Universe"]
    header(u, 1, UNIVERSE_COLS, [9, 9, 9, 11, 9, 16, 15, 7])
    for r in range(2, 301):
        row = DEFAULT_UNIVERSE[r - 2] if r - 2 < len(DEFAULT_UNIVERSE) else [None] * 8
        for j, v in enumerate(row, start=1):
            c = u.cell(r, j, v if v != "" else None)
            c.font, c.alignment = blue, Alignment(horizontal="left")
    for col, opts in (("B", "EQUITY,FI"), ("C", "AUTO,BASKET,PROXY"),
                      ("G", "SPDR,INVESCO,YAHOO,MANUAL"), ("H", "Y,N")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        dv.add(f"{col}2:{col}300")
        u.add_data_validation(dv)
    u.freeze_panes = "A2"

    # Holdings
    h = ws["Holdings"]
    header(h, 1, ["ETF", "Symbol", "Name", "Weight", "Chg %", "Contribution", "Priced Wt", "Source", "As Of"],
           [7, 9, 30, 9, 9, 11, 10, 9, 11], right_from=4)
    h.freeze_panes = "A2"

    # Charts
    ch = ws["Charts"]
    ch["A1"], ch["B1"], ch["A2"], ch["B2"] = "ETF", "XLK", "Days", 126
    ch["A1"].font = ch["A2"].font = bold
    ch["B1"].font = ch["B2"].font = blue
    ch.column_dimensions["A"].width = 8
    dv = DataValidation(type="list", formula1="=Universe!$A$2:$A$300")
    dv.add("B1")
    ch.add_data_validation(dv)

    # Config
    cf = ws["Config"]
    header(cf, 1, ["Setting", "Value"], [24, 10], right_from=2)
    for i, (label, _, default, name) in enumerate(CONFIG, start=2):
        cf.cell(i, 1, label).font = base
        v = cf.cell(i, 2, default)
        v.font, v.alignment = blue, Alignment(horizontal="right")
        if name:
            wb.defined_names[name] = DefinedName(name, attr_text=f"Config!$B${i}")
    for label, opts in (("Live", "Y,N"), ("Anchor (CLOSE/NAV)", "CLOSE,NAV"), ("Log intraday", "Y,N")):
        dv = DataValidation(type="list", formula1=f'"{opts}"')
        dv.add(f"B{2 + [c[0] for c in CONFIG].index(label)}")
        cf.add_data_validation(dv)

    # Feed
    fd = ws["Feed"]
    header(fd, 1, FEED_COLS + ["", "Session Left"], [12] + [10] * (len(FEED_COLS) + 1), right_from=2)
    fd["R2"].number_format = "0%"
    wb.defined_names["SessionLeft"] = DefinedName("SessionLeft", attr_text="Feed!$R$2")
    fd.freeze_panes = "B2"
    wb.save(BOOK)


def _ignore_formula_flags(path: Path, sheet_index: int, sqref: str) -> None:
    """Turn off Excel's green 'inconsistent formula' triangles on the Dashboard (openpyxl can't)."""
    import shutil
    import zipfile
    tmp = path.with_suffix(".tmp")
    name = f"xl/worksheets/sheet{sheet_index}.xml"
    tag = (f'<ignoredErrors><ignoredError sqref="{sqref}" formula="1" formulaRange="1" '
           f'numberStoredAsText="1" evalError="1" emptyCellReference="1"/></ignoredErrors>')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == name:
                xml = data.decode("utf-8")
                cut = min([i for i in (xml.find("<drawing"), xml.find("<legacyDrawing"), xml.find("<extLst"))
                           if i >= 0] or [xml.rfind("</worksheet>")])
                data = (xml[:cut] + tag + xml[cut:]).encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


# ------------------------------------------------------------------ open / read
def open_book() -> xw.Book:
    for app in xw.apps:
        for b in app.books:
            if b.name == BOOK.name:
                return b
    if not BOOK.exists():
        build_workbook()
    book = xw.Book(str(BOOK))
    try:  # cells written later pick up the workbook's default font
        st = book.api.styles["Normal"].font_object
        st.name.set(FONT)
        st.size.set(10)
    except Exception:
        pass
    return book


def read_settings(book: xw.Book) -> Settings:
    s = Settings()
    vals = book.sheets["Config"][f"A2:B{len(CONFIG) + 1}"].value
    by_label = {str(r[0]).strip(): r[1] for r in vals if r[0]}
    for label, fld, default, _ in CONFIG:
        v = by_label.get(label, default)
        v = default if v in (None, "") else v
        cur = getattr(s, fld)
        if isinstance(cur, bool):
            v = str(v).strip().upper().startswith("Y")
        elif isinstance(cur, int):
            v = int(float(v))
        elif isinstance(cur, float):
            v = float(v)
        else:
            v = str(v).strip()
        setattr(s, fld, v)
    return s


def _lev(x) -> float | None:
    try:
        x = float(str(x).lower().replace("x", ""))
        return x or None
    except (TypeError, ValueError):
        return None


def read_universe(book: xw.Book) -> list[dict]:
    rows = book.sheets["Universe"].range("A2:H300").value
    out, seen = [], set()
    for i, r in enumerate(rows, start=2):
        etf = yahoo_symbol(r[0]) if r[0] else ""
        if not etf or etf in seen or str(r[7] or "Y").strip().upper().startswith("N"):
            continue
        seen.add(etf)
        cls = str(r[1] or "").strip().upper()
        model = str(r[2] or "AUTO").strip().upper()
        out.append({
            "row": i, "etf": etf,
            "class": "FI" if cls in ("FI", "FIXED INCOME", "BOND", "BONDS") else "EQUITY" if cls else "",
            "model": model if model in ("AUTO", "BASKET", "PROXY") else "AUTO",
            "hedge": yahoo_symbol(r[3]) if r[3] else "",
            "leverage": _lev(r[4]),
            "peers": [yahoo_symbol(p) for p in str(r[5] or "").replace(";", ",").split(",") if p.strip()],
            "source": str(r[6] or "").strip(),
        })
    return out


FI_WORDS = re.compile(r"bond|treasur|government|corporate|muni|inflation|credit|fixed|mortgage", re.I)
NOT_TICKERS = {"ETF", "DAILY", "BULL", "BEAR", "SHARES", "TRUST", "FUND", "LONG", "SHORT", "INVERSE", "INC",
               "THE", "AND", "INDEX", "DIREXION", "GRANITE", "TRADR", "LEVERAGED", "USD", "ETFS", "AMPLIFY"}


def autofill_universe(book: xw.Book, uni: list[dict]) -> None:
    """A row with only a ticker gets its class, model, underlying/leverage and holdings source filled in."""
    sh = book.sheets["Universe"]
    for u in uni:
        if u["class"]:
            continue
        try:
            info = yf.Ticker(u["etf"]).info
        except Exception:
            continue
        name = f'{info.get("longName") or ""} {info.get("shortName") or ""}'
        u["class"] = "FI" if FI_WORDS.search(f'{name} {info.get("category") or ""}') else "EQUITY"
        lev = re.search(r"(-?\d+(?:\.\d+)?)\s*[xX]\b", name)
        if lev and not u["leverage"]:
            u["leverage"] = float(lev.group(1)) * (-1 if re.search(r"bear|short|inverse", name, re.I) else 1)
            for tok in re.findall(r"\b[A-Z]{1,5}\b", name):
                if tok in NOT_TICKERS or tok == u["etf"]:
                    continue
                try:
                    if yf.Ticker(tok).info.get("quoteType") == "EQUITY":
                        u["hedge"] = tok
                        break
                except Exception:
                    pass
        if u["leverage"] or u["class"] == "FI":
            u["model"] = "PROXY"
            u["hedge"] = u["hedge"] or ("IEF" if u["class"] == "FI" else "SPY")
        elif not u["source"]:
            for src in ("SPDR", "INVESCO", "YAHOO"):
                try:
                    if len(fetch_holdings(u["etf"], src)):
                        u["source"] = src
                        break
                except Exception:
                    pass
            u["model"] = "BASKET" if u["source"] else "PROXY"
            u["hedge"] = u["hedge"] or ("" if u["source"] else "SPY")
        sh.range((u["row"], 2)).value = [u["class"], u["model"], u["hedge"] or None, u["leverage"],
                                         ",".join(u["peers"]) or None, u["source"] or None, "Y"]
        print(f"  {u['etf']}: {u['class']} / {u['model']} / {u['hedge'] or u['source']}"
              + (f" / {u['leverage']:g}x" if u["leverage"] else ""))


def read_holdings(book: xw.Book) -> pd.DataFrame:
    sh = book.sheets["Holdings"]
    last = sh.range("A" + str(sh.cells.last_cell.row)).end("up").row
    cols = ["etf", "symbol", "name", "weight", "source", "as_of"]
    if last < 2:
        return pd.DataFrame(columns=cols)
    a = sh.range(f"A2:D{last}").value
    b = sh.range(f"H2:I{last}").value
    df = pd.DataFrame([x + y for x, y in zip(a, b)], columns=cols)
    df = df[df["etf"].notna() & df["symbol"].notna()].copy()
    df["etf"] = df["etf"].map(yahoo_symbol)
    df["symbol"] = df["symbol"].map(yahoo_symbol)
    df["weight"] = pd.to_numeric(df["weight"], errors="coerce")
    df = df.dropna(subset=["weight"])
    pct = df.groupby("etf")["weight"].transform("sum") > 1.5   # typed as 5 instead of 5%
    df.loc[pct, "weight"] /= 100
    return df


def write_holdings(book: xw.Book, df: pd.DataFrame) -> None:
    sh = book.sheets["Holdings"]
    sh.range("A2:I20000").clear_contents()
    if not len(df):
        return
    n = len(df) + 1
    sh["A2"].value = [[r.etf, r.symbol, r.name, r.weight] for r in df.itertuples()]
    sh["H2"].value = [[r.source, r.as_of] for r in df.itertuples()]
    chg = (f'=IFERROR(LET(m,MATCH($B{{r}},Feed!$A:$A,0),p,INDEX(Feed!${F["Last"]}:${F["Last"]},m),'
           f'c,INDEX(Feed!${F["Prev Close"]}:${F["Prev Close"]},m),IF(AND(p>0,c>0),p/c-1,"")),"")')
    sh[f"E2:G{n}"].formula = [[chg.format(r=r), f"=IF(ISNUMBER(E{r}),D{r}*E{r},0)", f"=IF(ISNUMBER(E{r}),D{r},0)"]
                        for r in range(2, n + 1)]
    sh[f"D2:G{n}"].number_format = "0.00%"
    sh[f"F2:F{n}"].number_format = "0.000%"


def load_holdings(book: xw.Book, uni: list[dict], force: bool) -> None:
    cur = read_holdings(book)
    frames, changed = [cur], False
    for u in uni:
        src = u["source"]
        if not src or src.upper() == "MANUAL" or (not force and (cur["etf"] == u["etf"]).any()):
            continue
        try:
            df = fetch_holdings(u["etf"], src)
        except Exception as e:
            print(f"  holdings {u['etf']} ({src}): failed - {e}")
            continue
        print(f"  holdings {u['etf']}: {len(df)} names from {src if '//' not in src else 'URL'}")
        frames[0] = frames[0][frames[0]["etf"] != u["etf"]]
        frames.append(df.assign(etf=u["etf"], source=src.upper() if "//" not in src else "URL",
                                as_of=dt.date.today().isoformat()))
        changed = True
    if changed:
        out = pd.concat(frames, ignore_index=True)[["etf", "symbol", "name", "weight", "source", "as_of"]]
        write_holdings(book, out.astype(object).where(out.notna(), None))


# ------------------------------------------------------------------ engine (data feed)
class Engine:
    def __init__(self):
        self.yahoo, self.nasdaq, self.alpaca = Yahoo(), Nasdaq(), Alpaca.from_env()
        self.meta: dict[str, dict] = {}
        self.meta_t = self.retry_t = 0.0
        self.incomplete: set[str] = set()
        self.fallback = False
        self.key = None
        self.models: dict[str, Model] = {}
        self.universe: list[dict] = []
        self.feed_keys: list[str] = []

    def history(self, symbols: list[str], period: str) -> tuple[pd.DataFrame, pd.DataFrame]:
        """Daily history, cached on disk per day so restarts and new tickers only fetch what's missing."""
        CACHE.mkdir(exist_ok=True)
        f = CACHE / f"history_{period}_{dt.date.today():%Y-%m-%d}.pkl"
        close, vol = pd.read_pickle(f) if f.exists() else (pd.DataFrame(), pd.DataFrame())
        need = [x for x in symbols if x not in close.columns]
        if need:
            print(f"loading {period} history for {len(need)} symbols...")
            c, v = self.yahoo.daily_history(need, period)
            if len(c):
                close = pd.concat([close, c], axis=1) if len(close) else c
                vol = pd.concat([vol, v], axis=1) if len(vol) else v
                pd.to_pickle((close, vol), f)
                for old in CACHE.glob("history_*.pkl"):
                    if old != f:
                        old.unlink()
        return close, vol

    def maybe_rebuild(self, s: Settings, universe: list[dict], holdings: pd.DataFrame) -> bool:
        self.s = s
        key = hashlib.md5(repr(([{k: v for k, v in u.items() if k != "row"} for u in universe],
                                holdings[["etf", "symbol", "weight"]].values.tolist(), s.max_basket_names,
                                s.history_period, s.beta_lookback, s.z_lookback, s.vol_lookback,
                                s.holdings_shown, dt.date.today())).encode()).hexdigest()
        if key == self.key and not (self.incomplete and time.time() > self.retry_t):
            return False
        self.universe = universe
        self.etfs = [u["etf"] for u in universe]
        extra = set()
        for u in universe:
            extra |= {u["hedge"], *u["peers"]} - {""}
            if u["model"] == "BASKET" or (u["model"] == "AUTO" and u["class"] == "EQUITY"):
                w = holdings[holdings["etf"] == u["etf"]].nlargest(s.max_basket_names, "weight")
                extra |= set(w["symbol"])
        self.others = sorted(extra - set(self.etfs))
        self.close, self.vol = self.history(self.etfs + self.others, s.history_period)
        core = set(self.etfs) | {u["hedge"] for u in universe} | {p for u in universe for p in u["peers"]}
        self.incomplete = (core - {""}) - set(self.close.columns)
        self.retry_t = time.time() + 300
        self.models = {u["etf"]: build_model(u, holdings, self.close, s) for u in universe}
        done = self.close[self.close.index.date < dt.date.today()]
        self.sigma = done.pct_change().tail(s.vol_lookback).std()
        self.basket_sigma = {e: float(m.und_hist.tail(s.vol_lookback).std())
                             for e, m in self.models.items() if m.kind == "BASKET" and m.und_hist is not None}
        self.feed_keys = self.etfs + self.others + [f"{e} BASKET" for e in self.basket_sigma]
        self.key = key
        return True

    def shown_under(self, m: Model) -> list[tuple[str, str]]:
        """(role, key) rows displayed under an ETF."""
        if m.kind == "BASKET":
            top = m.weights.sort_values(ascending=False).head(self.s.holdings_shown).index
            return [("Basket", f"{m.etf} BASKET")] + [("Holding", x) for x in top]
        return [("Underlying", m.hedge)] if m.hedge else []

    def tick(self) -> list[list]:
        if time.time() - self.meta_t > 600:          # NAV / yield change once a day
            meta = self.yahoo.etf_meta(self.etfs)
            self.meta.update({k: v for k, v in meta.items() if v})
            self.meta_t = time.time()
        try:
            yq, self.fallback = self.yahoo.quotes(self.etfs + self.others), False
        except Exception:                             # Yahoo throttling: Nasdaq, 20 names per call
            etf_like = set(self.etfs) | {u["hedge"] for u in self.universe} | {p for u in self.universe for p in u["peers"]}
            yq, self.fallback = self.nasdaq.batch(self.etfs + self.others, etf_like), True
        shown = {k for m in self.models.values() for role, k in self.shown_under(m) if role != "Basket"}
        live_syms = list(dict.fromkeys(self.etfs + sorted(shown)))
        stocks = {x for x in live_syms if yq.get(x, {}).get("type") == "EQUITY"}
        live = self.alpaca.quotes(self.etfs + self.others) if self.alpaca else self.nasdaq.quotes(live_syms, stocks)
        rows = []
        for key in self.feed_keys:
            if key.endswith(" BASKET"):
                rows.append([key] + [None] * 10 + [self.basket_sigma.get(key.split()[0])] + [None] * 4)
                continue
            q = dict(yq.get(key, {}))
            q.update({k: v for k, v in live.get(key, {}).items() if v is not None and not
                      (isinstance(v, float) and np.isnan(v))})
            meta = self.meta.get(key, {})
            bid, ask, last = q.get("bid", np.nan), q.get("ask", np.nan), q.get("last", np.nan)
            ok = np.isfinite(bid) and np.isfinite(ask) and ask >= bid > 0
            if ok and np.isfinite(last):   # a quote far from the last trade is stale: drop it
                mid = (bid + ask) / 2
                ok = abs(mid / last - 1) <= max(1.5 * (ask - bid) / mid, 0.0010)
            prev = q.get("prev_close", np.nan)
            if not (np.isfinite(prev) and prev > 0) and key in self.close and self.close[key].notna().any():
                prev = self.close[key].dropna().iloc[-1]
            m = self.models.get(key)
            rows.append([key, last, bid if ok else None, ask if ok else None,
                         q.get("bid_size") if ok else None, q.get("ask_size") if ok else None, prev,
                         q.get("volume"), q.get("avg_volume"), meta.get("nav"), meta.get("yield"),
                         self.sigma.get(key), q.get("time"),
                         m.beta if m else None, m.resid_mean if m else None, m.resid_std if m else None])
        return [[_clean(v) for v in r] for r in rows]


def _clean(v):
    if isinstance(v, (float, np.floating)):
        return None if not np.isfinite(v) else float(v)
    return v


ENGINE = Engine()


# ------------------------------------------------------------------ Dashboard formulas
def dashboard_rows() -> list[tuple[list, str]]:
    """Each ETF row followed by its underlying rows, as formulas. Returns (row, role)."""
    out: list[tuple[list, str]] = []
    groups = [[u for u in ENGINE.universe if u["class"] != "FI"], [u for u in ENGINE.universe if u["class"] == "FI"]]
    r = FIRST
    for members in groups:
        if out and members:
            out.append(([""] * (len(DASH) + 5), "Blank"))
            r += 1
        for u in members:
            m = ENGINE.models[u["etf"]]
            under = ENGINE.shown_under(m)
            etf_row = r
            out.append((_formula_row(r, "ETF", u["etf"], u["etf"], u["class"], etf_row, bool(under)), "ETF"))
            r += 1
            for role, key in under:
                label = "Basket" if role == "Basket" else key
                out.append((_formula_row(r, role, label, key, u["class"], etf_row, True), role))
                r += 1
    return out


def _formula_row(r: int, role: str, label: str, key: str, cls: str, etf_row: int, has_under: bool) -> list:
    def idx(col: str) -> str:
        return f'INDEX(Feed!${F[col]}:${F[col]},${HIDDEN["Row"]}{r})'

    def look(col: str) -> str:
        return f'=IFERROR(IF({idx(col)}="","",{idx(col)}),"")'

    c = {h: f"{C[h]}{r}" for h in C}
    row = {h: "" for h in C}
    row["Ticker"] = label
    market = role in ("ETF", "Underlying", "Holding")
    if market:
        row.update({
            "Last": look("Last"), "Bid": look("Bid"), "Ask": look("Ask"),
            "Bid Sz": look("Bid Sz"), "Ask Sz": look("Ask Sz"),
            "Spread bps": (f'=IF(AND(N({c["Bid"]})>0,N({c["Ask"]})>0),'
                           f'({c["Ask"]}-{c["Bid"]})/(({c["Ask"]}+{c["Bid"]})/2)*10000,"")'),
            "Chg %": f'=IF(N({c["Last"]})>0,IFERROR({c["Last"]}/{idx("Prev Close")}-1,""),"")',
            "Volume": look("Volume"), "Avg Vol": look("Avg Vol"),
            "Rel Vol": f'=IF(AND(N({c["Volume"]})>0,N({c["Avg Vol"]})>0),{c["Volume"]}/{c["Avg Vol"]},"")',
            "Time": look("Time"),
        })
    if role == "Basket":
        e = f"$A${etf_row}"
        row["β / Wt"] = (f'=IFERROR(SUMIFS(Holdings!$G:$G,Holdings!$A:$A,{e})'
                         f'/SUMIFS(Holdings!$D:$D,Holdings!$A:$A,{e}),"")')
        row["Chg %"] = (f'=IFERROR(SUMIFS(Holdings!$F:$F,Holdings!$A:$A,{e})'
                        f'/SUMIFS(Holdings!$G:$G,Holdings!$A:$A,{e}),"")')
    if role == "Holding":
        row["β / Wt"] = f"=SUMIFS(Holdings!$D:$D,Holdings!$A:$A,$A${etf_row},Holdings!$B:$B,{c['Ticker']})"
    row["σ 1D"] = look("Sigma")
    row["Move σ"] = f'=IF(AND(ISNUMBER({c["Chg %"]}),N({c["σ 1D"]})>0),{c["Chg %"]}/{c["σ 1D"]},"")'
    z = c["Move σ"]
    for k in (1, 2, 3):
        row[f"P>{k}σ"] = (f'=IF(ISNUMBER({z}),IF(SessionLeft<=0,--(ABS({z})>{k}),'
                          f'NORM.S.DIST((-{k}-{z})/SQRT(SessionLeft),TRUE)'
                          f'+1-NORM.S.DIST(({k}-{z})/SQRT(SessionLeft),TRUE)),"")')
    if role == "ETF":
        leg_row = etf_row + 1
        leg = f"UPPER($A${leg_row})" if has_under else '"HEDGE"'
        cost = f'IF(${HIDDEN["Class"]}{r}="FI",CostFI,CostEq)'
        edge = f'IF(${HIDDEN["Class"]}{r}="FI",EdgeFI,EdgeEq)'
        prev = idx("Prev Close")
        fv, px = c["Fair Value"], c["Last"]
        row.update({
            "β / Wt": look("Beta"),
            "NAV T-1": look("NAV"),
            "Prem T-1 bps": f'=IFERROR(IF(N({c["NAV T-1"]})>0,({prev}/{c["NAV T-1"]}-1)*10000,""),"")',
            "Und Move": f'=IFERROR({c["β / Wt"]}*${C["Chg %"]}${leg_row},"")' if has_under else "",
            "Fair Value": (f'=IF(ISNUMBER({c["Und Move"]}),IF(AND(Anchor="NAV",N({c["NAV T-1"]})>0),'
                           f'{c["NAV T-1"]},{prev})*(1+{c["Und Move"]}),"")'),
            "vs FV bps": (f'=IF(ISNUMBER({fv}),(IF(AND(N({c["Bid"]})>0,N({c["Ask"]})>0),'
                          f'({c["Bid"]}+{c["Ask"]})/2,{px})/{fv}-1)*10000,"")'),
            "Sell Edge": f'=IF(ISNUMBER({fv}),(IF(N({c["Bid"]})>0,{c["Bid"]},{px})/{fv}-1)*10000-{cost},"")',
            "Buy Edge": f'=IF(ISNUMBER({fv}),(1-IF(N({c["Ask"]})>0,{c["Ask"]},{px})/{fv})*10000-{cost},"")',
            "Resid Z": f'=IFERROR(({c["Chg %"]}-{c["Und Move"]}-{idx("Resid Mean")})/{idx("Resid SD")},"")',
            "Yield": look("Yield"),
            "Signal": (f'=IF(NOT(ISNUMBER({fv})),"",IF((NOW()-N({c["Time"]}))*1440>StaleMin,"STALE",'
                       f'IF(AND(${HIDDEN["Class"]}{r}="EQUITY",ABS(N({c["Prem T-1 bps"]}))>ExDivFlag),"EX-DIV?",'
                       f'IF({c["Sell Edge"]}>{edge},"SELL ETF / BUY "&{leg},'
                       f'IF({c["Buy Edge"]}>{edge},"BUY ETF / SELL "&{leg},"")))))'),
        })
    hidden = [key, role, cls, f'=IFERROR(MATCH(${HIDDEN["Key"]}{r},Feed!$A:$A,0),"")']
    return [row[h] for h, _, _ in DASH] + [""] + hidden


class Dashboard:
    def __init__(self, book: xw.Book):
        self.book, self.sh = book, book.sheets["Dashboard"]
        self.layout = None
        self.etf_rows: dict[str, int] = {}

    def lay_out(self) -> None:
        rows = dashboard_rows()
        layout = [r[0] for r, _ in rows]
        if layout == self.layout:
            return
        self.sh.range(f"A{FIRST}:{HIDDEN['Row']}{MAX_ROW}").clear_contents()
        if rows:  # xlwings on Mac needs the full target range for a 2D formula write
            self.sh.range((FIRST, 1), (FIRST + len(rows) - 1, len(rows[0][0]))).formula = [r for r, _ in rows]
        self.etf_rows = {r[0]: FIRST + i for i, (r, role) in enumerate(rows) if role == "ETF"}
        self.layout = layout

    def read_etfs(self) -> list[dict]:
        if not self.etf_rows:
            return []
        vals = self.sh.range(f"A{FIRST}:{C['Time']}{max(self.etf_rows.values())}").value
        heads = [h for h, _, _ in DASH]
        return [dict(zip(heads, vals[r - FIRST])) for r in self.etf_rows.values()]


def write_feed(book: xw.Book, rows: list[list], n_prev: int) -> None:
    sh = book.sheets["Feed"]
    if n_prev > len(rows):
        sh.range(f"A{len(rows) + 2}:P{n_prev + 1}").clear_contents()
    sh["A2"].value = rows
    sh["R2"].value = session_left()
    if n_prev != len(rows):
        n = len(rows) + 1
        for a, b, fmt in (("B", "D", "#,##0.00"), ("E", "F", "#,##0"), ("G", "G", "#,##0.00"),
                          ("H", "I", "#,##0"), ("J", "J", "#,##0.00"), ("K", "L", "0.00%"),
                          ("M", "M", "hh:mm:ss"), ("N", "N", "0.00"), ("O", "P", "0.000%")):
            sh[f"{a}2:{b}{n}"].number_format = fmt


def log_live(rows: list[dict]) -> None:
    LOGS.mkdir(exist_ok=True)
    f = LOGS / f"live_{dt.date.today():%Y-%m-%d}.csv"
    keep = ["Ticker", "Bid", "Ask", "Last", "Fair Value", "vs FV bps", "Spread bps", "Sell Edge", "Buy Edge",
            "Volume", "Signal"]
    df = pd.DataFrame([{k: (r.get(k) if r.get(k) != "" else None) for k in keep} for r in rows])
    df.insert(0, "time", dt.datetime.now().replace(microsecond=0))
    if f.exists() and f.open().readline().strip() != ",".join(df.columns):
        f.rename(f.with_suffix(".old.csv"))           # columns changed: start a new file
    df.to_csv(f, mode="a", header=not f.exists(), index=False)


# ------------------------------------------------------------------ charts
NAVY, ORANGE, GREYC, LIGHT = "#1f3864", "#c55a11", "#7f7f7f", "#8faadc"


def draw_charts(book: xw.Book, etf: str, days: int) -> None:
    sh = book.sheets["Charts"]
    m = ENGINE.models.get(etf)
    if m is None:
        return
    hist = history_frame(m, ENGINE.close, ENGINE.vol, days)
    und = "Basket" if m.kind == "BASKET" else f"{m.beta:.2f}x {m.hedge}"
    rets = ENGINE.close[etf].pct_change().dropna() if etf in ENGINE.close else pd.Series(dtype=float)
    fig, ax = plt.subplots(2, 3, figsize=(17, 8.5))

    a = ax[0, 0]
    if len(hist):
        a.plot(hist.index, hist["ETF"], label=etf, lw=1.6, color=NAVY)
        a.plot(hist.index, hist["Underlying"], label=und, lw=1.2, color=ORANGE)
        for p in m.peers:
            if p in hist:
                a.plot(hist.index, hist[p], label=p, lw=0.8, alpha=0.7)
        a.legend(fontsize=8)
    a.set_title(f"{etf} vs {und}")

    a = ax[0, 1]
    if len(hist):
        sp = hist["Spread (bps)"]
        mu, sd = sp.rolling(20).mean(), sp.rolling(20).std()
        a.fill_between(sp.index, mu - 2 * sd, mu + 2 * sd, color=GREYC, alpha=0.15, lw=0)
        a.plot(sp.index, sp, color=NAVY, lw=1.3)
        a.plot(mu.index, mu, color=GREYC, lw=0.8)
        a.axhline(0, color="black", lw=0.5)
    a.set_title("Tracking spread (bps)")

    a = ax[0, 2]
    if len(rets) > 20:
        zs = rets.tail(max(days, 250)) / rets.tail(ENGINE.s.vol_lookback).std()
        a.hist(zs, bins=np.arange(-6, 6.25, 0.25), color=LIGHT, density=True)
        x = np.linspace(-5, 5, 200)
        a.plot(x, np.exp(-x ** 2 / 2) / np.sqrt(2 * np.pi), color=NAVY, lw=1)
        for k in (-3, -2, -1, 1, 2, 3):
            a.axvline(k, color=GREYC, lw=0.5, ls=":")
        freq = [f"{(zs.abs() > k).mean():.1%}" for k in (1, 2, 3)]
        a.set_title(f"Daily moves (σ)   >1σ {freq[0]}   >2σ {freq[1]}   >3σ {freq[2]}")
        a.set_xlim(-5, 5)

    a = ax[1, 0]
    if len(hist) and hist["Volume"].notna().any():
        v = hist["Volume"] / 1e6
        a.bar(v.index, v, color=LIGHT, width=1.0)
        a.plot(v.index, v.rolling(20).mean(), color=NAVY, lw=1.1)
    a.set_title("Volume (mm)")

    a = ax[1, 1]
    if len(rets):
        vol = (rets.rolling(20).std() * 100).tail(days)
        a.plot(vol.index, vol, color=NAVY, lw=1.3)
    a.set_title("20d σ (daily %)")

    a = ax[1, 2]
    f = LOGS / f"live_{dt.date.today():%Y-%m-%d}.csv"
    lg = pd.read_csv(f, parse_dates=["time"]) if f.exists() else pd.DataFrame()
    lg = lg[lg["Ticker"] == etf] if len(lg) else lg
    if len(lg) > 1:
        a.plot(lg["time"], lg["vs FV bps"], color=NAVY, lw=1.3, label="vs FV")
        a.plot(lg["time"], lg["Sell Edge"], color="#c00000", lw=0.8, label="sell edge")
        a.plot(lg["time"], lg["Buy Edge"], color="#548235", lw=0.8, label="buy edge")
        a.plot(lg["time"], lg["Spread bps"], color=GREYC, lw=0.8, ls="--", label="spread")
        a.axhline(0, color="black", lw=0.5)
        a.legend(fontsize=8)
        a.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M"))
    a.set_title("Intraday (bps)")

    for a in (ax[0, 0], ax[0, 1], ax[1, 0], ax[1, 1]):
        a.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=3, maxticks=6))
        a.xaxis.set_major_formatter(mdates.DateFormatter("%b %y"))
    fig.tight_layout()
    sh.pictures.add(fig, name="etf_charts", update=True, anchor=sh["A4"])
    plt.close(fig)

    hs = book.sheets["Chart Data"]
    hs.range("A1:Z3000").clear_contents()
    if len(hist):
        out = hist.round(4)
        out.index.name = "Date"
        hs["A1"].options(index=True, header=True).value = out


# ------------------------------------------------------------------ loop
def refresh(book: xw.Book, dash: Dashboard, state: dict) -> float:
    s = read_settings(book)
    if not s.live:
        book.sheets["Dashboard"]["A1"].value = f"Paused {dt.datetime.now():%H:%M:%S}"
        return s.refresh_seconds
    uni = read_universe(book)
    if any(not u["class"] for u in uni):
        autofill_universe(book, uni)
    new = [u for u in uni if u["source"] and u["source"].upper() != "MANUAL" and u["etf"] not in state["tried"]]
    if new:
        state["tried"] |= {u["etf"] for u in new}
        load_holdings(book, uni, force=False)
    holdings = read_holdings(book)
    n_prev = len(ENGINE.feed_keys)
    ENGINE.maybe_rebuild(s, uni, holdings)
    t0 = time.time()
    rows = ENGINE.tick()
    write_feed(book, rows, n_prev)
    dash.lay_out()
    took = time.time() - t0
    src = "Alpaca" if ENGINE.alpaca else "Nasdaq"
    book.sheets["Dashboard"]["A1"].value = (f"{dt.datetime.now():%H:%M:%S}   {src} / "
                                            f"{'Nasdaq (Yahoo throttled)' if ENGINE.fallback else 'Yahoo'}")
    etf_vals = dash.read_etfs()
    if s.log_live:
        log_live(etf_vals)
    etf, days = book.sheets["Charts"]["B1:B2"].value
    etf, days = yahoo_symbol(etf or ""), int(days or 126)
    if (etf, days, ENGINE.key) != state.get("chart") or time.time() - state.get("chart_t", 0) > s.chart_seconds:
        draw_charts(book, etf, days)
        state["chart"], state["chart_t"] = (etf, days, ENGINE.key), time.time()
    sig = sum(1 for r in etf_vals if str(r.get("Signal") or "").startswith(("SELL", "BUY")))
    print(f"{dt.datetime.now():%H:%M:%S}  {len(etf_vals)} ETFs, {sig} signal(s), fetch {took:.1f}s")
    return s.refresh_seconds


def redraw_if_changed(book: xw.Book, state: dict, force: bool = False) -> None:
    """Redraw the charts when Charts!B1/B2 change, also outside market hours."""
    etf, days = book.sheets["Charts"]["B1:B2"].value
    etf, days = yahoo_symbol(etf or ""), int(days or 126)
    if not force and (etf, days, ENGINE.key) == state.get("chart"):
        return
    if ENGINE.key is None or etf not in ENGINE.models:
        ENGINE.maybe_rebuild(read_settings(book), read_universe(book), read_holdings(book))
    draw_charts(book, etf, days)
    state["chart"], state["chart_t"] = (etf, days, ENGINE.key), time.time()
    print(f"{dt.datetime.now():%H:%M:%S}  charts: {etf}, {days} days")


def market_phase(now: dt.datetime) -> str:
    """'pre' (holdings time), 'open' (stream), or 'closed', in New York time."""
    t = now.astimezone(ET)
    if t.weekday() >= 5:
        return "closed"
    hm = t.hour * 60 + t.minute
    return "open" if 9 * 60 + 25 <= hm <= 16 * 60 + 5 else "pre" if 8 * 60 + 30 <= hm < 9 * 60 + 25 else "closed"


def holdings_are_current(book: xw.Book) -> bool:
    h = read_holdings(book)
    return len(h) > 0 and (h["as_of"].astype(str) == dt.date.today().isoformat()).all()


def run(once: bool, auto: bool = False) -> None:
    book = open_book()
    book.activate()
    dash = Dashboard(book)
    state: dict = {"tried": set()}
    print(f"ETF arb tracker on {BOOK.name}{' (auto: market hours only)' if auto else ''}. Ctrl+C to stop.")
    while True:
        t0 = time.time()
        try:
            if auto:
                phase = market_phase(dt.datetime.now(ET))
                if phase != "closed" and not holdings_are_current(book):
                    print(f"{dt.datetime.now():%H:%M:%S}  refreshing holdings")
                    load_holdings(book, read_universe(book), force=True)
                    book.save()
                if phase != "open":
                    if state.get("phase") == "open":
                        book.save()                      # keep the closing snapshot
                    state["phase"] = phase
                    book.sheets["Dashboard"]["A1"].value = f"{dt.datetime.now():%H:%M}   market closed"
                    redraw_if_changed(book, state)
                    time.sleep(5)
                    continue
                state["phase"] = phase
            wait = refresh(book, dash, state)
        except KeyboardInterrupt:
            raise
        except Exception as e:  # Excel busy (cell in edit mode), network hiccup, rate limit ...
            limited = "Rate limit" in str(e) or "Too Many Requests" in str(e)
            wait = 90 if limited else 10
            msg = f"{dt.datetime.now():%H:%M:%S}   {'Yahoo rate limit' if limited else 'refresh failed'}, retry in {wait}s"
            print(f"{msg}: {str(e).splitlines()[0]}")
            try:
                book.sheets["Dashboard"]["A1"].value = msg
            except Exception:                       # workbook was closed: reopen it
                try:
                    book = open_book()
                    dash = Dashboard(book)
                except Exception:
                    pass
            if os.environ.get("ETF_ARB_DEBUG"):
                traceback.print_exc()
        if once:
            book.save()
            return
        time.sleep(max(1.0, wait - (time.time() - t0)))


def main() -> None:
    p = argparse.ArgumentParser(description="Live ETF arbitrage monitor in Excel")
    p.add_argument("command", nargs="?", default="run", choices=["run", "auto", "holdings", "setup", "chart"])
    p.add_argument("ticker", nargs="?", help="chart: ETF to draw (default: Charts!B1)")
    p.add_argument("--once", action="store_true", help="refresh once and exit")
    p.add_argument("--force", action="store_true", help="setup: overwrite the existing workbook")
    a = p.parse_args()
    if a.command == "setup":
        if BOOK.exists() and not a.force:
            sys.exit(f"{BOOK.name} exists; use --force to rebuild it")
        keep = None
        if BOOK.exists():  # carry the user's tickers and settings into the new workbook
            old = open_book()
            keep = (old.sheets["Universe"]["A2:H300"].value,
                    {r[0]: r[1] for r in old.sheets["Config"]["A2:B60"].value if r[0]})
        for app in xw.apps:
            for name in [b.name for b in app.books]:
                if name == BOOK.name:
                    app.books[name].close()
        BOOK.unlink(missing_ok=True)
        build_workbook()
        book = open_book()
        if keep:
            book.sheets["Universe"]["A2:H300"].value = keep[0]
            cfg = book.sheets["Config"]
            for i, (label, *_ ) in enumerate(CONFIG, start=2):
                if keep[1].get(label) not in (None, ""):
                    cfg.range((i, 2)).value = keep[1][label]
        load_holdings(book, read_universe(book), force=True)
        run(once=True)                 # writes the formulas
        book.close()
        # Excel only keeps "ignore error" ranges over cells that already hold formulas
        _ignore_formula_flags(BOOK, sheet_index=1, sqref=f"A{FIRST}:{HIDDEN['Row']}{MAX_ROW}")
        xw.Book(str(BOOK)).sheets["Dashboard"].activate()
        print(f"created {BOOK}")
    elif a.command == "chart":
        book = open_book()
        if a.ticker:
            book.sheets["Charts"]["B1"].value = a.ticker.upper()
        redraw_if_changed(book, {}, force=True)
        book.sheets["Charts"].activate()
    elif a.command == "holdings":
        book = open_book()
        load_holdings(book, read_universe(book), force=True)
        book.save()
    else:
        try:
            run(a.once, auto=a.command == "auto")
        except KeyboardInterrupt:
            print("stopped")


if __name__ == "__main__":
    main()
