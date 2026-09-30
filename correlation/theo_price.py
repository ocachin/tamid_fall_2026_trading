"""Theoretical price of the Direxion Bear 1X ETFs from beta and from NAV, against where they traded.

A -1x daily fund resets its exposure every close, so during session d:

    R_U        = (underlying now + its dividend if ex today) / underlying close(d-1) - 1
    Theo (β)   = ETF close(d-1) × (1 + β × R_U + carry) - distribution if the fund is ex today
    Theo (NAV) = NAV(d-1)       × (1 -     R_U + carry) - distribution

β is -1 by contract; the workbook also prices with the estimated daily β to show it adds nothing.
Direxion's holdings file gives NAV(d-1) exactly (swaps = -100% of net assets, the rest money-market
funds); its product page gives the last 30 NAVs rounded to the cent. Python writes prices, NAVs and
distributions; every return, beta, theo, premium and error statistic is an Excel formula.

    .venv/bin/python correlation/theo_price.py                     # last 7 sessions
    .venv/bin/python correlation/theo_price.py --replay correlation/cache/bars_2026-09-23.pkl \\
                                              --funds correlation/cache/funds_2026-09-30.pkl
"""
from __future__ import annotations

import argparse
import datetime as dt
import logging
import os
import pickle
from pathlib import Path

import numpy as np
import pandas as pd
import yfinance as yf
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import direxion
import intraday_corr_resolved as res
from intraday_corr_resolved import FONT, header, yahoo

logging.getLogger("yfinance").setLevel(logging.ERROR)

HERE = Path(__file__).resolve().parent
BOOK = HERE / "Theo_Price.xlsx"
CARRY_BPS = 1.8          # net income per calendar day: money-market yield + swap financing - fees
NP = 10                  # columns per pair on the minute sheet
NS = 6                   # columns per pair on the Sessions sheet
NYSE_HOLIDAYS = ["2026-01-01", "2026-01-19", "2026-02-16", "2026-04-03", "2026-05-25", "2026-06-19",
                 "2026-07-03", "2026-09-07", "2026-11-26", "2026-12-25", "2027-01-01", "2027-01-18",
                 "2027-02-15", "2027-03-26", "2027-05-31", "2027-06-18", "2027-07-05", "2027-09-06",
                 "2027-11-25", "2027-12-24"]
SESSION = pd.offsets.CustomBusinessDay(holidays=NYSE_HOLIDAYS)


def next_session(t: pd.Timestamp, idx: pd.DatetimeIndex) -> pd.Timestamp:
    """The session after t: from Yahoo's daily bars when they reach that far, else the NYSE calendar."""
    later = idx[idx > t]
    return later[0] if len(later) else t + SESSION


# ------------------------------------------------------------------ data
def cached_funds() -> list[dict]:
    """Earlier fund snapshots, newest first."""
    out = []
    for p in sorted(res.CACHE.glob("funds_*.pkl"), reverse=True):
        with open(p, "rb") as f:
            out.append(pickle.load(f))
    return out


def fetch_funds(pairs: list) -> dict:
    """Distributions, NAVs and today's holdings from Direxion; underlying dividends from Yahoo.

    A fund Direxion won't serve falls back to the newest snapshot in cache/ that has it."""
    out = {"dist": {}, "nav": {}, "holdings": {}, "und_div": {}}
    older = None
    for e, u in pairs:
        try:
            out["holdings"][e] = direxion.holdings(e)
            out["dist"][e] = direxion.distributions(e)
            out["nav"][e] = direxion.nav_history(e)
        except Exception as err:                               # 403, timeout, page layout change
            older = cached_funds() if older is None else older
            snap = next((o for o in older if e in o["dist"]), None)
            if snap is None:
                raise SystemExit(f"Direxion failed for {e} ({err}) and no cached snapshot has it; "
                                 "rerun later or pass --funds with an earlier cache/funds_*.pkl")
            print(f"  WARNING: Direxion failed for {e} ({err}); using the newest cached snapshot")
            for k in ("holdings", "dist", "nav"):
                out[k][e] = snap[k][e]
        d = yf.Ticker(yahoo(u)).dividends
        if d is None:
            raise SystemExit(f"Yahoo returned no dividend data for {yahoo(u)} (offline?)")
        d.index = pd.DatetimeIndex(d.index).tz_localize(None).normalize()
        out["und_div"][u] = d
        nav = out["nav"][e]
        print(f"  {e}: {len(out['dist'][e])} distributions, {len(nav)} NAVs to {nav.index.max().date()}, "
              f"exact NAV {direxion.exact_nav(out['holdings'][e]):.5f}")
    return out


def exact_navs(funds: dict, pairs: list, daily_close: pd.DataFrame) -> dict:
    """Exact NAVs from the saved holdings files, dated against the closes being priced (plus any an
    older funds snapshot carried, for a machine without those files)."""
    old = funds.get("nav_exact", {})
    return {e: direxion.archived_navs(e, daily_close[yahoo(u)]).combine_first(old.get(e, pd.Series(dtype=float)))
            for e, u in pairs}


def sessions_table(days: list, pairs: list, daily: pd.DataFrame, funds: dict, exact_nav: dict) -> pd.DataFrame:
    """Per session: calendar days of carry until the next session (a Friday's NAV accrues the weekend),
    then per pair the prior closes, distribution, dividend and NAV(d-1)."""
    idx = daily.index
    rows = []
    for d0 in days:
        t = pd.Timestamp(d0)
        prev = idx[idx < t][-1]
        row = {"date": t, "carry_days": (next_session(t, idx) - t).days}
        for e, u in pairs:
            if daily.loc[prev, [yahoo(e), yahoo(u)]].isna().any():
                raise SystemExit(f"Yahoo has no {prev.date()} close for {e} or {u}; cannot anchor {t.date()}")
            exact, cent = exact_nav[e], funds["nav"][e]
            nav, src = (exact[prev], "exact") if prev in exact.index else \
                       (cent[prev], "cent") if prev in cent.index else (np.nan, "")
            row |= {(e, "etf_prev"): daily.at[prev, yahoo(e)], (e, "dist"): funds["dist"][e].get(t, 0.0),
                    (e, "und_prev"): daily.at[prev, yahoo(u)], (e, "und_div"): funds["und_div"][u].get(t, 0.0),
                    (e, "nav_prev"): nav, (e, "nav_src"): src}
        rows.append(row)
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ workbook
def build(close: pd.DataFrame, daily_close: pd.DataFrame, daily_adj: pd.DataFrame, days: list, pairs: list,
          funds: dict, path: Path) -> None:
    exact_nav = exact_navs(funds, pairs, daily_close)
    ses = sessions_table(days, pairs, daily_close, funds, exact_nav)
    wb = Workbook()
    wb.remove(wb.active)
    base, bold = Font(name=FONT, size=10), Font(name=FONT, size=10, bold=True)
    blue = Font(name=FONT, size=10, color="0000FF")
    names = ["Summary", "Beta", "By Day", "Today", "Holdings", "NAV", "Sessions", "Theo 1m", "RD", "Config"]
    ws = {n: wb.create_sheet(n) for n in names}
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False

    # ---- Config: the two assumptions, as inputs
    cfg = ws["Config"]
    for r, (lab, v, fmt) in enumerate((("Carry (bps per calendar day)", CARRY_BPS, "0.0"), ("Exposure", -1, "0.00")), 2):
        cfg.cell(r, 1, lab).font = base
        c = cfg.cell(r, 2, v)
        c.font, c.number_format = blue, fmt
    header(cfg, 1, 1, "Input", bold, left=True, width=28)
    header(cfg, 1, 2, "Value", bold, width=10)
    CARRY, EXPO = "Config!$B$2", "Config!$B$3"

    # ---- Sessions: prices and cash flows each session's theo is anchored on
    s = ws["Sessions"]
    header(s, 1, 1, "Session", bold, left=True, width=12)
    header(s, 1, 2, "Carry days", bold, width=8)
    for k, (e, u) in enumerate(pairs):
        for j, lab in enumerate((f"{e} close d-1", f"{e} distribution", f"{u} close d-1", f"{u} dividend",
                                 f"{e} NAV d-1", f"{e} NAV source")):
            header(s, 1, 3 + NS * k + j, lab, bold, width=11)
    for i in range(2, 2 + len(ses)):
        rec = ses.iloc[i - 2]
        c = s.cell(i, 1, rec["date"].to_pydatetime())
        c.font, c.number_format = base, "yyyy-mm-dd"
        s.cell(i, 2, int(rec["carry_days"])).font = base
        for k, (e, u) in enumerate(pairs):
            for j, key in enumerate(("etf_prev", "dist", "und_prev", "und_div", "nav_prev", "nav_src")):
                v = rec[(e, key)]
                c = s.cell(i, 3 + NS * k + j, None if isinstance(v, float) and np.isnan(v) else
                           v if isinstance(v, str) else float(v))
                c.font, c.number_format = base, "0.00000" if key in ("dist", "und_div", "nav_prev") else "0.00"
    s.freeze_panes = "C2"
    last_s = 1 + len(ses)

    def sref(k: int, j: int) -> str:
        col = get_column_letter(3 + NS * k + j)
        return f"Sessions!${col}$2:${col}${last_s}"

    # ---- Theo 1m: prices (values), returns since the prior close, three theos and premiums (formulas)
    t = ws["Theo 1m"]
    header(t, 1, 1, "Time", bold, left=True, width=17)
    header(t, 1, 2, "Session row", bold, width=8)
    labels = ("{e}", "{u}", "{e} return since close", "{u} return since close", "Theo β=-1",
              "Theo β est.", "Theo NAV", "Premium β=-1 (bps)", "Premium β est. (bps)", "Premium NAV (bps)")
    for k, (e, u) in enumerate(pairs):
        for j, lab in enumerate(labels):
            header(t, 1, 3 + NP * k + j, lab.format(e=e, u=u), bold, width=11)
    grid = close.index
    n = len(grid)
    last_t = n + 1
    beta_row = {k: 2 + k for k in range(len(pairs))}          # Beta sheet row per pair
    for i, ts in enumerate(grid, start=2):
        c = t.cell(i, 1, ts.tz_localize(None).to_pydatetime())
        c.font, c.number_format = base, "yyyy-mm-dd hh:mm"
        c = t.cell(i, 2, f"=MATCH(INT($A{i}),Sessions!$A$2:$A${last_s},0)")
        c.font = base
        for k, (e, u) in enumerate(pairs):
            cols = [get_column_letter(3 + NP * k + j) for j in range(NP)]
            E, U, RE, RU, T1, T2, T3, P1, P2, P3 = cols
            pe, pu = close.at[ts, yahoo(e)], close.at[ts, yahoo(u)]
            for col, v in ((E, pe), (U, pu)):
                c = t[f"{col}{i}"]
                c.value = None if np.isnan(v) else float(v)
                c.font, c.number_format = base, "0.00"
            ix = lambda j: f"INDEX({sref(k, j)},$B{i})"                     # noqa: E731
            ru = f"({U}{i}+{ix(3)})/{ix(2)}-1"
            carry = f"{CARRY}/10000*INDEX(Sessions!$B$2:$B${last_s},$B{i})"
            f = {RE: f'=IF(ISNUMBER({E}{i}),({E}{i}+{ix(1)})/{ix(0)}-1,"")',
                 RU: f'=IF(AND(ISNUMBER({E}{i}),ISNUMBER({U}{i})),{ru},"")',
                 T1: f'=IF(ISNUMBER({U}{i}),{ix(0)}*(1+{EXPO}*({ru})+{carry})-{ix(1)},"")',
                 T2: f'=IF(ISNUMBER({U}{i}),{ix(0)}*(1+Beta!$C${beta_row[k]}*({ru})+{carry})-{ix(1)},"")',
                 T3: f'=IF(AND(ISNUMBER({U}{i}),ISNUMBER({ix(4)})),{ix(4)}*(1+{EXPO}*({ru})+{carry})-{ix(1)},"")'}
            for col, th in ((P1, T1), (P2, T2), (P3, T3)):
                f[col] = f'=IF(AND(ISNUMBER({E}{i}),ISNUMBER({th}{i})),({E}{i}/{th}{i}-1)*10000,"")'
            for col, v in f.items():
                c = t[f"{col}{i}"]
                c.value = v
                c.font = base
                c.number_format = "0.0000%" if col in (RE, RU) else "0.0000" if col in (T1, T2, T3) else "0.0"
    t.freeze_panes = "C2"

    def tcol(k: int, j: int, lo: int = 2, hi: int = last_t) -> str:
        col = get_column_letter(3 + NP * k + j)
        return f"'Theo 1m'!${col}${lo}:${col}${hi}"

    # ---- RD: dividend-adjusted daily closes for the daily beta
    last_d = res.write_daily(ws["RD"], daily_adj, pairs, bold, base)
    first_session = pd.Timestamp(days[0])
    prior_hi = 1 + int((daily_adj.index < first_session).sum())      # last RD row before the window

    # ---- Beta: contractual, daily (out of sample: the year before the window), and since the prior close
    b = ws["Beta"]
    heads = ["Pair", "Contract", "β daily (1y before window)", "SE", "R² daily", "β since close (window)",
             "R² since close", "n minutes"]
    for j, h in enumerate(heads, start=1):
        header(b, 1, j, h, bold, left=j == 1, width=12 if j > 1 else 13)
    for k, (e, u) in enumerate(pairs):
        r = 2 + k
        da, db = get_column_letter(4 + 4 * k), get_column_letter(5 + 4 * k)
        y, x = f"RD!${da}$2:${da}${prior_hi}", f"RD!${db}$2:${db}${prior_hi}"
        vals = {1: f"{e}/{u}", 2: f"={EXPO}", 3: f"=SLOPE({y},{x})", 4: f"=STEYX({y},{x})/SQRT(DEVSQ({x}))",
                5: f"=RSQ({y},{x})", 6: f"=SLOPE({tcol(k, 2)},{tcol(k, 3)})", 7: f"=RSQ({tcol(k, 2)},{tcol(k, 3)})",
                8: f"=COUNT({tcol(k, 2)})"}
        for j, v in vals.items():
            c = b.cell(r, j, v)
            c.font, c.number_format = base, "#,##0" if j == 8 else "0.000"

    # ---- Summary: how far the ETF traded from each theo (only minutes in which it traded)
    sm = ws["Summary"]
    heads = ["Pair", "Tick (bps)"]
    for th in ("β=-1", "β est.", "NAV"):
        heads += [f"Mean {th}", f"SD {th}", f"RMSE {th}"]
    heads += ["n"]
    for j, h in enumerate(heads, start=1):
        header(sm, 1, j, h, bold, left=j == 1, width=13 if j == 1 else 10)
    for k, (e, u) in enumerate(pairs):
        r = 2 + k
        sm.cell(r, 1, f"{e}/{u}").font = base
        c = sm.cell(r, 2, f"=0.01/AVERAGE({tcol(k, 0)})*10000")
        c.font, c.number_format = base, "0.0"
        for m, j in enumerate((7, 8, 9)):
            p = tcol(k, j)
            for off, fx in enumerate((f"=AVERAGE({p})", f"=STDEV({p})", f"=SQRT(SUMSQ({p})/COUNT({p}))")):
                c = sm.cell(r, 3 + 3 * m + off, fx)
                c.font, c.number_format = base, "0.0"
        c = sm.cell(r, 12, f"=COUNT({tcol(k, 9)})")
        c.font, c.number_format = base, "#,##0"
    for col in ("E", "H", "K"):
        sm.conditional_formatting.add(f"{col}2:{col}{1 + len(pairs)}", ColorScaleRule(
            start_type="min", start_color="63BE7B", end_type="max", end_color="F8696B"))

    # ---- By Day: the same statistics for each session (the ex-dividend day included)
    bd = ws["By Day"]
    heads = ["Session", "ETF", "Underlying", "Distribution", "Mean β=-1", "Mean β est.", "Mean NAV",
             "RMSE β=-1", "RMSE β est.", "RMSE NAV", "n"]
    for j, h in enumerate(heads, start=1):
        header(bd, 1, j, h, bold, left=j <= 3, width=11)
    for si, d0 in enumerate(days):
        lo = 2 + int((close.index.date < d0).sum())
        hi = lo + int((close.index.date == d0).sum()) - 1
        for k, (e, u) in enumerate(pairs):
            r = 2 + si * len(pairs) + k
            c = bd.cell(r, 1, d0)
            c.font, c.number_format = base, "yyyy-mm-dd"
            bd.cell(r, 2, e).font = base
            bd.cell(r, 3, u).font = base
            c = bd.cell(r, 4, f"=INDEX({sref(k, 1)},{si + 1})")
            c.font, c.number_format = base, "0.00000"
            for m, j in enumerate((7, 8, 9)):
                p = tcol(k, j, lo, hi)
                for col, fx in ((5 + m, f'=IFERROR(AVERAGE({p}),"")'), (8 + m, f'=IFERROR(SQRT(SUMSQ({p})/COUNT({p})),"")')):
                    c = bd.cell(r, col, fx)
                    c.font, c.number_format = base, "0.0"
            c = bd.cell(r, 11, f"=COUNT({tcol(k, 9, lo, hi)})")
            c.font, c.number_format = base, "#,##0"
    bd.auto_filter.ref = f"A1:K{1 + len(days) * len(pairs)}"
    bd.freeze_panes = "D2"

    # ---- Holdings: today's Direxion files, and each fund's NAV and exposure worked out from them
    h = ws["Holdings"]
    cols = ["Fund", "Description", "Kind", "Counterparty", "Shares", "Price", "Market value", "% of net assets"]
    for j, lab in enumerate(cols, start=1):
        header(h, 1, j, lab, bold, left=j <= 4, width=(8, 34, 13, 18, 14, 10, 16, 12)[j - 1])
    r = 2
    for e, _u in pairs:
        for p in funds["holdings"][e].itertuples():
            for j, v in enumerate((p.fund, p.description, p.kind, p.counterparty, p.shares, p.price,
                                   p.market_value, p.pct_net_assets / 100), start=1):
                c = h.cell(r, j, v)
                c.font = base
                c.number_format = {5: "#,##0", 6: "#,##0.00", 7: "#,##0", 8: "0.00%"}.get(j, "General")
            r += 1
    last_h = r - 1
    fx = 10
    for j, lab in enumerate(("Fund", "Shares outstanding", "Net assets", "NAV", "Swap exposure",
                             "Money market", "Swap counterparties", "Close described"), start=fx):
        header(h, 1, j, lab, bold, left=j == fx, width=13)
    rng = lambda col: f"${col}$2:${col}${last_h}"                           # noqa: E731
    for k, (e, u) in enumerate(pairs):
        r = 2 + k
        hd = funds["holdings"][e]
        d = direxion.nav_date(hd, daily_close[yahoo(u)])
        mm = f'SUMIFS({rng("G")},{rng("A")},$J{r},{rng("C")},"money market")'
        vals = [e, hd.attrs["shares_outstanding"],
                f'={mm}/SUMIFS({rng("H")},{rng("A")},$J{r},{rng("C")},"money market")', f"=L{r}/K{r}",
                f'=SUMIFS({rng("G")},{rng("A")},$J{r},{rng("C")},"swap")/L{r}', f"={mm}/L{r}",
                f'=COUNTIFS({rng("A")},$J{r},{rng("C")},"swap")', d.to_pydatetime() if d is not None else None]
        for j, v in enumerate(vals, start=fx):
            c = h.cell(r, j, v)
            c.font = base
            c.number_format = {11: "#,##0", 12: "#,##0", 13: "0.00000", 14: "0.00%", 15: "0.00%",
                               17: "yyyy-mm-dd"}.get(j, "General")
    h.freeze_panes = "B2"

    # ---- NAV: Direxion's NAVs next to the official closes; the premium is where the close printed
    nv = ws["NAV"]
    dates = sorted(set().union(*[funds["nav"][e].index for e, _ in pairs]))
    dates = [d for d in dates if d in daily_close.index]
    header(nv, 1, 1, "Date", bold, left=True, width=12)
    for k, (e, _u) in enumerate(pairs):
        for j, lab in enumerate((f"{e} NAV", f"{e} close", f"{e} close vs NAV (bps)")):
            header(nv, 1, 2 + 3 * k + j, lab, bold, width=11)
    for i, d in enumerate(dates, start=2):
        c = nv.cell(i, 1, d.to_pydatetime())
        c.font, c.number_format = base, "yyyy-mm-dd"
        for k, (e, _u) in enumerate(pairs):
            a, cl, pr = (get_column_letter(2 + 3 * k + j) for j in range(3))
            navs = exact_nav[e].combine_first(funds["nav"][e])
            for col, v in ((a, navs.get(d, np.nan)), (cl, daily_close.at[d, yahoo(e)])):
                c = nv[f"{col}{i}"]
                c.value = None if pd.isna(v) else float(v)
                c.font, c.number_format = base, "0.0000"
            c = nv[f"{pr}{i}"]
            c.value = f'=IF(AND(ISNUMBER({a}{i}),ISNUMBER({cl}{i})),({cl}{i}/{a}{i}-1)*10000,"")'
            c.font, c.number_format = base, "0.0"
    nv.freeze_panes = "B2"

    # ---- Today: the next session's theo from the latest close and NAV. Type an underlying price
    td = ws["Today"]
    last_close = daily_close.index[daily_close.index <= pd.Timestamp(days[-1])][-1]
    nxt = next_session(last_close, daily_close.index)
    carry_days = (next_session(nxt, daily_close.index) - nxt).days
    heads = ["ETF", "Underlying", "Close date", "ETF close", "Underlying close", "NAV", "NAV source",
             "Distribution", "Dividend", "Carry days", "Underlying price", "ETF price", "Theo β=-1", "Theo NAV",
             "Premium to theo NAV (bps)"]
    for j, lab in enumerate(heads, start=1):
        header(td, 1, j, lab, bold, left=j <= 2, width=11)
    for k, (e, u) in enumerate(pairs):
        r = 2 + k
        exact, cent = exact_nav[e], funds["nav"][e]
        nav, src = (exact[last_close], "exact") if last_close in exact.index else \
                   (cent[last_close], "cent") if last_close in cent.index else (None, "not yet published")
        ec, uc = float(daily_close.at[last_close, yahoo(e)]), float(daily_close.at[last_close, yahoo(u)])
        vals = [e, u, last_close.to_pydatetime(), ec, uc, nav, src, funds["dist"][e].get(nxt, 0.0),
                funds["und_div"][u].get(nxt, 0.0), carry_days, uc, ec,
                f"=D{r}*(1+{EXPO}*((K{r}+I{r})/E{r}-1)+{CARRY}/10000*J{r})-H{r}",
                f'=IF(ISNUMBER(F{r}),F{r}*(1+{EXPO}*((K{r}+I{r})/E{r}-1)+{CARRY}/10000*J{r})-H{r},"")',
                f'=IF(ISNUMBER(N{r}),(L{r}/N{r}-1)*10000,"")']
        for j, v in enumerate(vals, start=1):
            c = td.cell(r, j, v)
            c.font = blue if j in (11, 12) else base
            c.number_format = {3: "yyyy-mm-dd", 4: "0.00", 5: "0.00", 6: "0.00000", 8: "0.00000", 9: "0.00000",
                               11: "0.00", 12: "0.00", 13: "0.0000", 14: "0.0000", 15: "0.0"}.get(j, "General")

    tmp = path.with_name(f"~{path.name}")
    wb.save(tmp)
    os.replace(tmp, path)


def main() -> None:
    p = argparse.ArgumentParser(description="Theo price of the Bear 1X ETFs from beta and from NAV")
    p.add_argument("--sessions", type=int, default=7)
    p.add_argument("--replay", type=Path, help="bars saved by intraday_corr_resolved.py")
    p.add_argument("--funds", type=Path, help="Direxion/dividend data saved by an earlier run")
    p.add_argument("--out", type=Path, default=BOOK)
    p.add_argument("--open", action="store_true")
    args = p.parse_args()

    if args.replay:
        with open(args.replay, "rb") as f:
            raw = pickle.load(f)
    else:
        print("downloading 1m and daily bars ...")
        raw = res.fetch(args.sessions)
        res.snapshot(raw)
    data, days, pairs = res.prepare(raw, args.sessions, partial=False)
    daily_close = raw["daily_close"].copy()
    daily_close.index = pd.DatetimeIndex(daily_close.index).tz_localize(None).normalize()
    if args.funds:
        with open(args.funds, "rb") as f:
            funds = pickle.load(f)
        lacking = [e for e, _ in pairs if e not in funds["dist"]]
        if lacking:
            raise SystemExit(f"{args.funds.name} has no data for {', '.join(lacking)}; rerun without --funds")
    else:
        print("fetching Direxion holdings, NAVs and distributions ...")
        funds = fetch_funds(pairs)
        res.CACHE.mkdir(exist_ok=True)
        snap = res.CACHE / f"funds_{dt.date.today()}.pkl"
        with open(snap, "wb") as f:
            pickle.dump(res.portable(funds), f)
        print(f"  saved to {snap.relative_to(HERE.parent)} (rerun with --funds to reuse it)")
    print(f"  sessions {days[0]} .. {days[-1]} ({len(days)})")
    build(data["close"], daily_close, data["daily"], days, pairs, funds, args.out)
    print(f"wrote {args.out}")
    if args.open:
        import xlwings as xw
        xw.Book(str(args.out)).sheets["Summary"].activate()


if __name__ == "__main__":
    main()
