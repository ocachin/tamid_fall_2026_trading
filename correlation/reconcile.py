"""Why `etf intra script.py` and correlation/intraday_corr.py disagree, on one set of bars.

Rebuilds both published tables from the same 1-minute Yahoo bars, then splits the gap between them into
five method choices with an exact Shapley decomposition: each factor's step, averaged over all 120 orders
in which the five can be switched. Prints the resolved method's figures alongside.

    statistic     price levels (Haas)                   -> returns (original)
    missing       drop minutes without a trade (Haas)   -> forward fill inside the session (original)
    bars          Yahoo bars, labelled at bin start     -> 1m resampled with label/closed='right' (original)
    window        7 sessions (Haas)                     -> last 5 sessions (original)
    aggregation   mean of per-session ρ (Haas)          -> one ρ over all sessions pooled (original)

    .venv/bin/python correlation/reconcile.py --replay correlation/cache/bars_2026-09-23.pkl
"""
from __future__ import annotations

import argparse
import itertools
import math
import pickle
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

import intraday_corr_resolved as res

HERE = Path(__file__).resolve().parent
HAAS_BOOK = HERE.parent / "New ETF Data.xlsx"
ORIG_BOOK = HERE / "Correlation_Intraday.xlsx"
INTERVALS = [1, 5, 15]
HAAS = dict(statistic="levels", missing="dropna", bars="native", window=7, aggregation="per_session")
ORIG = dict(statistic="returns", missing="ffill", bars="right", window=5, aggregation="pooled")
FACTORS = list(HAAS)


def bar_prices(close: pd.DataFrame, m: int, bars: str) -> pd.DataFrame:
    """Close on the bar grid; NaN = no trade in the bar. 'native' reproduces Yahoo's own 5m/15m bars."""
    if m == 1:
        return close
    side = "left" if bars == "native" else "right"   # intraday_corr.returns() resamples label/closed='right'
    parts = [g.resample(f"{m}min", label=side, closed=side).last() for _, g in close.groupby(close.index.date)]
    return pd.concat(parts)                    # forward filling is the 'missing' factor's job (same result)


def pair_values(px: pd.DataFrame, e: str, u: str, cfg: dict) -> pd.DataFrame:
    """Columns x (ETF), y (underlying), day: levels or within-session returns of one pair."""
    out = []
    for day, g in px[[e, u]].groupby(px.index.date):
        g = (g.ffill() if cfg["missing"] == "ffill" else g).dropna()
        if cfg["statistic"] == "returns":
            g = g.pct_change().iloc[1:]         # never across the overnight gap
        out.append(pd.DataFrame({"x": g[e].to_numpy(), "y": g[u].to_numpy(), "day": day}))
    return pd.concat(out, ignore_index=True)


def corr(close: pd.DataFrame, days: list, cfg: dict, pairs: list, per_day: bool = False):
    win = days if cfg["window"] == 7 else days[-5:]
    table, rows = pd.DataFrame(index=[f"{e}/{u}" for e, u in pairs], columns=INTERVALS, dtype=float), []
    for m in INTERVALS:
        px = bar_prices(close, m, cfg["bars"])
        for e, u in pairs:
            v = pair_values(px, res.yahoo(e), res.yahoo(u), cfg)
            v = v[v["day"].isin(win)]
            if cfg["aggregation"] == "pooled":
                table.loc[f"{e}/{u}", m] = v["x"].corr(v["y"])
            else:
                rs = [(d, g["x"].corr(g["y"])) for d, g in v.groupby("day") if len(g) >= 3]
                rows += [(d, e, u, m, r) for d, r in rs]
                table.loc[f"{e}/{u}", m] = np.mean([r for _, r in rs])
    if per_day:
        return table, pd.DataFrame(rows, columns=["day", "ETF", "Underlying", "m", "r"])
    return table


def shapley(close: pd.DataFrame, days: list, pairs: list) -> dict[str, pd.DataFrame]:
    """Exact Shapley value of each factor's contribution to ORIG minus HAAS, per pair and interval."""
    value = {}
    for bits in itertools.product((0, 1), repeat=len(FACTORS)):
        cfg = {f: (ORIG if b else HAAS)[f] for f, b in zip(FACTORS, bits)}
        value[bits] = corr(close, days, cfg, pairs)
    n, out = len(FACTORS), {}
    for i, f in enumerate(FACTORS):
        phi = 0
        for bits in value:
            if bits[i]:
                continue
            s = sum(bits)
            w = math.factorial(s) * math.factorial(n - s - 1) / math.factorial(n)
            phi = phi + w * (value[bits[:i] + (1,) + bits[i + 1:]] - value[bits])
        out[f] = phi
    return out


def resolved(close: pd.DataFrame, days: list, pairs: list) -> tuple[pd.DataFrame, pd.DataFrame]:
    """The resolved method (refresh-time returns, last common trade per bin): per-session mean and pooled."""
    per, pooled = (pd.DataFrame(index=[f"{e}/{u}" for e, u in pairs], columns=INTERVALS, dtype=float) for _ in range(2))
    for m in INTERVALS:
        for e, u in pairs:
            b = res.sync_bars(close, e, u, m, days)
            prev = b[["e", "u"]].shift()
            ok = (b["sync"] == 1) & (b.index.date == pd.Series(b.index.date, index=b.index).shift().to_numpy()) & prev["e"].notna()
            r = (b[["e", "u"]] / prev - 1)[ok]
            pooled.loc[f"{e}/{u}", m] = r["e"].corr(r["u"])
            per.loc[f"{e}/{u}", m] = np.mean([g["e"].corr(g["u"]) for _, g in r.groupby(r.index.date)])
    return per, pooled


def check_published(close: pd.DataFrame, days: list, pairs: list) -> None:
    """Compare with the numbers in the two committed workbooks, where their sessions overlap ours."""
    if HAAS_BOOK.exists():
        _, mine = corr(close, days, HAAS, pairs, per_day=True)
        ws = openpyxl.load_workbook(HAAS_BOOK, data_only=True)["Master Summary"]
        pub = {(str(r[0]), r[1], m): v for r in ws.iter_rows(min_row=2, values_only=True)
               for m, v in zip(INTERVALS, r[3:6])}
        d = [abs(round(r.r, 4) - pub[(str(r.day), r.ETF, r.m)]) for r in mine.itertuples() if (str(r.day), r.ETF, r.m) in pub]
        print(f"Haas method reproduces {len(d)} cells of {HAAS_BOOK.name}: max |diff| {max(d) if d else float('nan'):.4f}")
    if ORIG_BOOK.exists():
        ws = openpyxl.load_workbook(ORIG_BOOK, data_only=True)["Chart"]
        rows = list(ws.iter_rows(min_row=2, max_row=9, values_only=True))
        cols = {m: rows[0].index(m) for m in INTERVALS}
        pub = {r[0].replace("^GSPC", "SPX"): r for r in rows[1:]}
        mine = corr(close, days, ORIG, pairs)
        d = [abs(mine.loc[p, m] - pub[p][cols[m]]) for p in mine.index for m in INTERVALS if p in pub]
        print(f"original method reproduces {len(d)} cells of {ORIG_BOOK.name} (Chart): max |diff| {max(d):.1e}"
              f"{'' if max(d) < 1e-9 else '  (the workbook was built from a different download window)'}")


def main() -> None:
    p = argparse.ArgumentParser(description="Reconcile the two intraday correlation scripts")
    p.add_argument("--replay", type=Path, required=True, help="a download saved by intraday_corr_resolved.py")
    args = p.parse_args()
    with open(args.replay, "rb") as f:
        raw = pickle.load(f)
    data, days, pairs = res.prepare(raw, 7, partial=False)
    close = data["close"]
    print(f"sessions {days[0]} .. {days[-1]}\n")
    check_published(close, days, pairs)

    pd.set_option("display.width", 160)
    fmt = lambda x: f"{x:+.3f}"                                        # noqa: E731
    haas, orig = corr(close, days, HAAS, pairs), corr(close, days, ORIG, pairs)
    per, pooled = resolved(close, days, pairs)
    side = pd.concat({"Haas": haas, "original": orig, "resolved (per session)": per, "resolved (pooled)": pooled}, axis=1)
    side.loc["mean"] = side.mean()
    print("\nρ by method (columns: minutes)\n" + side.to_string(float_format=lambda x: f"{x:.3f}"))

    phi = shapley(close, days, pairs)
    wf = pd.DataFrame({f: phi[f].mean() for f in FACTORS}).T
    wf.loc["total (original - Haas)"] = (orig - haas).mean()
    print("\nShapley split of the gap, mean over pairs (columns: minutes)\n" + wf.to_string(float_format=fmt))
    for m in INTERVALS:
        t = pd.DataFrame({f: phi[f][m] for f in FACTORS})
        t["total"] = orig[m] - haas[m]
        print(f"\nper pair, {m}m\n" + t.to_string(float_format=fmt))


if __name__ == "__main__":
    main()
