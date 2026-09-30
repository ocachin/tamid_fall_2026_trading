"""Intraday correlation and beta of the Direxion Bear 1X ETFs against their underlyings (resolved).

Reconciles `etf intra script.py` (per-session correlation of price levels) with
correlation/intraday_corr.py (pooled correlation of forward-filled returns). On the same Yahoo bars
the two differ only in method, so this version keeps what each got right:

  * returns, not price levels, and never across the overnight gap
  * a pair's return is measured only between minutes in which both legs traded (refresh time), so a
    thin ETF's stale price never scores a zero return against a moving underlying
  * 2-60 minute bars come from the 1-minute bars: each session starts from its 09:30 minute, then each
    bin (09:30, 09:35], (09:35, 09:40] ... is labelled by its last minute and holds both legs' prices at
    the last minute in it in which both traded. Every interval covers the same part of the day, the
    opening minutes included (the last bin of an interval that does not divide 390 is shorter)
  * every session on its own (By Day and the 1m/5m/15m matrices) and all sessions pooled (Summary)
  * daily returns from dividend-adjusted closes, so an ETF's distribution is not read as a price move

Python downloads prices and writes them. Every return, correlation, beta and R² in the workbook is an
Excel formula.

    .venv/bin/python correlation/intraday_corr_resolved.py                  # last 7 sessions
    .venv/bin/python correlation/intraday_corr_resolved.py --sessions 20    # Yahoo keeps ~30 days of 1m bars
    .venv/bin/python correlation/intraday_corr_resolved.py --replay correlation/cache/bars_2026-09-23.pkl
    .venv/bin/python correlation/intraday_corr_resolved.py --open           # ... and open it in Excel
"""
from __future__ import annotations

import argparse
import datetime as dt
import logging
import os
import pickle
from pathlib import Path

import numpy as np
import pandas as pd
import yfinance as yf
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

logging.getLogger("yfinance").setLevel(logging.ERROR)     # keeps "N Failed downloads: [...]" visible

HERE = Path(__file__).resolve().parent
BOOK = HERE / "Correlation_Resolved.xlsx"
CACHE = HERE / "cache"
NY = "America/New_York"
FONT, THIN = "Arial", Side(style="thin", color="000000")

PAIRS = [("AVS", "AVGO"), ("AAPD", "AAPL"), ("METD", "META"), ("MSFD", "MSFT"),
         ("SPDN", "SPX"), ("AMZD", "AMZN"), ("PLTD", "PLTR")]
YAHOO = {"SPX": "^GSPC"}
MINUTES = [1, 2, 3, 5, 10, 15, 30, 60]        # sampling intervals on the Epps sheet
BY_DAY = [1, 5, 15]                            # intervals shown per session
OPEN_MIN, SESSION_MIN = 9 * 60 + 30, 390
NCOL = 5                                       # columns per pair on a bar sheet


def yahoo(sym: str) -> str:
    return YAHOO.get(sym, sym)


# ------------------------------------------------------------------ data
def fetch(sessions: int) -> dict:
    """Raw 1-minute closes and volumes (NaN = no trade that minute) and 1 year of daily closes.

    Yahoo serves 1-minute bars for the last 30 days, at most 8 days per request."""
    syms = [yahoo(s) for p in PAIRS for s in p]
    end = dt.date.today() + dt.timedelta(days=1)
    a = end - dt.timedelta(days=min(29, sessions * 7 // 5 + 6))
    parts = []
    while a < end:
        b = min(a + dt.timedelta(days=7), end)
        d = yf.download(syms, start=a, end=b, interval="1m", auto_adjust=False, prepost=False,
                        progress=False, threads=True)
        if len(d):
            parts.append(d)
        a = b
    if not parts:
        raise SystemExit("Yahoo returned no 1-minute bars (offline or rate-limited?)")
    m = pd.concat(parts).sort_index()
    m = m[~m.index.duplicated()]
    daily = yf.download(syms, period="1y", interval="1d", auto_adjust=False, progress=False, threads=True)
    return {"close": m["Close"].reindex(columns=syms), "volume": m["Volume"].reindex(columns=syms),
            "daily_close": daily["Close"].reindex(columns=syms),
            "daily_adj": daily["Adj Close"].reindex(columns=syms)}


def to_ny(x: pd.DataFrame) -> pd.DataFrame:
    x = x.copy()
    x.index = (x.index.tz_localize("UTC") if x.index.tz is None else x.index).tz_convert(NY)
    return x


def in_progress(close: pd.DataFrame) -> dt.date | None:
    """Today's session if it may still be trading: its latest bar is under 20 minutes old. That also
    covers 13:00 closes, and leaves Yahoo a few minutes to settle the last bars after 16:00."""
    if close.empty:
        return None
    now, last = pd.Timestamp.now(tz=NY), close.index.max()
    return last.date() if last.date() == now.date() and now - last < pd.Timedelta(minutes=20) else None


def prepare(raw: dict, sessions: int, partial: bool) -> tuple[dict, list[dt.date], list[tuple[str, str]]]:
    """Regular-hours bars of the last `sessions` sessions; drops today's session while it is still open
    and any pair with a leg Yahoo returned nothing for, and warns about sessions a leg is missing from."""
    close = to_ny(raw["close"]).between_time("09:30", "15:59")
    volume = to_ny(raw["volume"]).reindex(close.index)
    days = sorted(set(close.index.date))
    if not partial and in_progress(close) in days:
        days = days[:-1]
    days = days[-sessions:]
    if not days:
        raise SystemExit("no complete regular-hours session in the download")
    keep = np.isin(close.index.date, days)
    close, volume = close[keep], volume[keep]

    daily = raw["daily_adj"].copy()
    daily.index = pd.DatetimeIndex(daily.index).tz_localize(None).normalize()
    daily = daily[daily.index.date <= days[-1]]          # no half-built bar for a session in progress

    pairs = []
    per_day = close.notna().groupby(close.index.date).sum()
    for e, u in PAIRS:
        ye, yu = yahoo(e), yahoo(u)
        missing = [s for s in (e, u) if close[yahoo(s)].notna().sum() == 0]
        if missing:
            print(f"  no data for {', '.join(missing)}: dropping {e}/{u}")
            continue
        pairs.append((e, u))
        for s, other in ((e, yu), (u, ye)):
            gap = [str(d) for d in days if per_day.at[d, yahoo(s)] == 0 and per_day.at[d, other] > 0]
            if gap:
                print(f"  WARNING: no {s} bars on {', '.join(gap)} (Yahoo download gap?); "
                      f"{e}/{u} uses the other sessions only")
        for s in (ye, yu):
            if daily[s].tail(len(days) + 1).isna().any():
                print(f"  WARNING: {s} has no daily close on some of these sessions (Yahoo gap?)")
    if not pairs:
        raise SystemExit("no pair has data for both legs")
    return {"close": close, "volume": volume, "daily": daily}, days, pairs


def trim(raw: dict, last: dt.date) -> dict:
    """The download without any bar after session `last`."""
    out = {}
    for k, x in raw.items():
        idx = to_ny(x).index.date if k in ("close", "volume") else pd.DatetimeIndex(x.index).date
        out[k] = x[idx <= last]
    return out


def portable(x):
    """pandas 3 pickles string labels in a form pandas 2 cannot read; store them as plain objects."""
    if isinstance(x, dict):
        return {k: portable(v) for k, v in x.items()}
    if isinstance(x, (pd.DataFrame, pd.Series)):
        x = x.copy()
        if x.index.dtype.kind in "OU" or str(x.index.dtype).startswith("str"):
            x.index = x.index.astype(object)
        if isinstance(x, pd.DataFrame):
            x.columns = x.columns.astype(object) if not isinstance(x.columns, pd.MultiIndex) else x.columns
            for c in x.columns:
                if str(x[c].dtype).startswith(("str", "string")):
                    x[c] = x[c].astype(object)
    return x


def save_snapshot(raw: dict, last: dt.date) -> Path:
    """Save the download up to session `last` to cache/bars_<last>.pkl, merged with any earlier file of
    that name, so a shorter download never replaces minutes Yahoo no longer serves."""
    CACHE.mkdir(exist_ok=True)
    path = CACHE / f"bars_{last}.pkl"
    new = trim(raw, last)
    if path.exists():
        with open(path, "rb") as f:
            old = pickle.load(f)
        new = {k: v.combine_first(old[k]) if k in old else v for k, v in new.items()}
    with open(path, "wb") as f:
        pickle.dump(portable(new), f)
    return path


def snapshot(raw: dict) -> None:
    """Keep a copy of a fresh download: complete sessions only (never one still trading, even with
    --partial), so that a later --replay cannot mistake a half-built session for a finished one."""
    close = to_ny(raw["close"]).between_time("09:30", "15:59")
    days = sorted(set(close.index.date) - {in_progress(close)})
    if days:
        path = save_snapshot(raw, days[-1])
        print(f"  saved the download to {os.path.relpath(path, HERE.parent)} (rerun with --replay to reuse it)")


def labels(m: int) -> list[int]:
    """Minutes after 09:30 of the rows of an m-minute grid: 0 (the 09:30 minute), then each bin's last
    minute; the session's last bin is capped at 390 (16:00) and holds minutes up to 15:59."""
    return sorted({min(-(-o // m) * m, SESSION_MIN) for o in range(SESSION_MIN)})


def sync_bars(close: pd.DataFrame, e: str, u: str, m: int, days: list[dt.date]) -> pd.DataFrame:
    """One pair on the m-minute grid of labels(m).

    In each bin: both legs' closes at the last minute in which both traded (refresh time), forward
    filled within the session so that the previous row always holds the previous synchronous price.
    'sync' = 1 where both legs traded in the bin, i.e. where a return ends."""
    both = close[[yahoo(e), yahoo(u)]].dropna()
    both.columns = ["e", "u"]
    t = both.index
    o = np.asarray(t.hour * 60 + t.minute - OPEN_MIN)
    lab = np.minimum(-(-o // m) * m, SESSION_MIN)
    last = both.groupby(t.normalize() + pd.to_timedelta(OPEN_MIN + lab, unit="min")).last()
    grid = pd.DatetimeIndex([pd.Timestamp(d).tz_localize(NY) + pd.Timedelta(minutes=OPEN_MIN + k)
                             for d in days for k in labels(m)])
    out = last.reindex(grid)
    out["sync"] = out["e"].notna().astype(int)
    out[["e", "u"]] = out[["e", "u"]].groupby(out.index.date).ffill()
    return out


# ------------------------------------------------------------------ workbook
def header(ws, row: int, col: int, text, bold, left: bool = False, width: float | None = None):
    c = ws.cell(row, col, text)
    c.font, c.border = bold, Border(bottom=THIN)
    c.alignment = Alignment(horizontal="left" if left else "right", wrap_text=True)
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


def pair_cols(k: int) -> tuple[str, str, str, str, str]:
    """ETF price, underlying price, both-traded flag, ETF return, underlying return."""
    return tuple(get_column_letter(2 + NCOL * k + j) for j in range(NCOL))


def write_bars(ws, bars: dict, pairs: list, bold, base) -> int:
    """Prices (values) and synchronous returns (formulas) for every pair on one grid. Returns last row."""
    header(ws, 1, 1, "Time", bold, left=True, width=17)
    for k, (e, u) in enumerate(pairs):
        for j, lab in enumerate((e, u, f"{e} both traded", f"{e} return", f"{u} return")):
            header(ws, 1, 2 + NCOL * k + j, lab, bold, width=11)
    grid = bars[pairs[0]].index
    arrays = {p: (bars[p]["e"].to_numpy(), bars[p]["u"].to_numpy(), bars[p]["sync"].to_numpy()) for p in pairs}
    for i, t in enumerate(grid, start=2):
        c = ws.cell(i, 1, t.tz_localize(None).to_pydatetime())
        c.font, c.number_format = base, "yyyy-mm-dd hh:mm"
        for k, p in enumerate(pairs):
            pe, pu, fl, re_, ru = pair_cols(k)
            ea, ua, sa = arrays[p]
            j0 = 2 + NCOL * k
            for j, v, fmt in ((0, ea[i - 2], "0.00"), (1, ua[i - 2], "0.00"), (2, int(sa[i - 2]), "0")):
                c = ws.cell(i, j0 + j, None if isinstance(v, float) and np.isnan(v) else float(v) if j < 2 else v)
                c.font, c.number_format = base, fmt
            if i == 2:
                continue
            same = f"INT($A{i})=INT($A{i - 1})"
            for j, px in ((3, pe), (4, pu)):
                c = ws.cell(i, j0 + j, f'=IF(AND({fl}{i}=1,{same},ISNUMBER({px}{i - 1})),{px}{i}/{px}{i - 1}-1,"")')
                c.font, c.number_format = base, "0.0000%"
    ws.freeze_panes = "B2"
    return len(grid) + 1


def write_daily(ws, daily: pd.DataFrame, pairs: list, bold, base) -> int:
    """Dividend-adjusted daily closes (values) and daily returns (formulas)."""
    header(ws, 1, 1, "Date", bold, left=True, width=12)
    for k, (e, u) in enumerate(pairs):
        for j, lab in enumerate((f"{e} adj close", f"{u} adj close", f"{e} return", f"{u} return")):
            header(ws, 1, 2 + 4 * k + j, lab, bold, width=11)
    for i, (t, row) in enumerate(daily.iterrows(), start=2):
        c = ws.cell(i, 1, t.to_pydatetime())
        c.font, c.number_format = base, "yyyy-mm-dd"
        for k, (e, u) in enumerate(pairs):
            pe, pu, re_, ru = (get_column_letter(2 + 4 * k + j) for j in range(4))
            for j, s in enumerate((e, u)):
                v = row[yahoo(s)]
                c = ws.cell(i, 2 + 4 * k + j, None if pd.isna(v) else float(v))
                c.font, c.number_format = base, "0.00"
            if i == 2:
                continue
            ok = f"AND(ISNUMBER({pe}{i}),ISNUMBER({pe}{i - 1}),ISNUMBER({pu}{i}),ISNUMBER({pu}{i - 1}))"
            for j, px in ((2, pe), (3, pu)):
                c = ws.cell(i, 2 + 4 * k + j, f'=IF({ok},{px}{i}/{px}{i - 1}-1,"")')
                c.font, c.number_format = base, "0.0000%"
    ws.freeze_panes = "B2"
    return len(daily) + 1


def heat(ws, rng: str) -> None:
    """Green at -1 (the ETF moves exactly against its underlying), red at 0."""
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="num", start_value=-1, start_color="63BE7B",
        mid_type="num", mid_value=-0.5, mid_color="FFEB84",
        end_type="num", end_value=0, end_color="F8696B"))


def build(data: dict, days: list, pairs: list, path: Path) -> None:
    close, volume, daily = data["close"], data["volume"], data["daily"]
    bars = {m: {p: sync_bars(close, *p, m, days) for p in pairs} for m in MINUTES}

    wb = Workbook()
    wb.remove(wb.active)
    base, bold = Font(name=FONT, size=10), Font(name=FONT, size=10, bold=True)
    summary = wb.create_sheet("Summary")
    byday = wb.create_sheet("By Day")
    matrices = {m: wb.create_sheet(f"{m}m Matrix") for m in BY_DAY}
    epps = wb.create_sheet("Epps")
    qual = wb.create_sheet("Data Quality")
    last = {m: write_bars(wb.create_sheet(f"R{m}"), bars[m], pairs, bold, base) for m in MINUTES}
    last_d = write_daily(wb.create_sheet("RD"), daily, pairs, bold, base)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False

    def rng(m: int | str, col: str, lo: int = 2, hi: int | None = None) -> str:
        sheet = f"R{m}" if m != "D" else "RD"
        hi = hi or (last[m] if m != "D" else last_d)
        return f"{sheet}!${col}${lo}:${col}${hi}"

    def daily_cols(k: int) -> tuple[str, str]:
        return get_column_letter(4 + 4 * k), get_column_letter(5 + 4 * k)

    # ---- Summary: all sessions pooled
    heads = ["Pair", "ETF", "Underlying", "ρ 1m", "ρ 5m", "ρ 15m", "ρ daily (1y)", "β 1m", "β 5m",
             "β 15m", "β daily (1y)", "R² 15m", "R² daily", "n 1m", "n 15m", "n daily", "ETF minutes traded"]
    for j, h in enumerate(heads, start=1):
        header(summary, 1, j, h, bold, left=j <= 3, width=11 if j <= 3 else 12)
    for k, (e, u) in enumerate(pairs):
        r = 2 + k
        _, _, fl, a, b = pair_cols(k)
        da, db = daily_cols(k)
        vals = {1: f"{e}/{u}", 2: e, 3: u,
                4: f"=CORREL({rng(1, a)},{rng(1, b)})", 5: f"=CORREL({rng(5, a)},{rng(5, b)})",
                6: f"=CORREL({rng(15, a)},{rng(15, b)})", 7: f"=CORREL({rng('D', da)},{rng('D', db)})",
                8: f"=SLOPE({rng(1, a)},{rng(1, b)})", 9: f"=SLOPE({rng(5, a)},{rng(5, b)})",
                10: f"=SLOPE({rng(15, a)},{rng(15, b)})", 11: f"=SLOPE({rng('D', da)},{rng('D', db)})",
                12: f"=RSQ({rng(15, a)},{rng(15, b)})", 13: f"=RSQ({rng('D', da)},{rng('D', db)})",
                14: f"=COUNT({rng(1, a)})", 15: f"=COUNT({rng(15, a)})", 16: f"=COUNT({rng('D', da)})",
                17: f"=COUNTIF({rng(1, fl)},1)/{len(close)}"}   # of the minutes Yahoo returned
        for j, v in vals.items():
            c = summary.cell(r, j, v)
            c.font = base
            c.number_format = "0.000" if 4 <= j <= 13 else "#,##0" if j <= 16 and j >= 14 else "0%" if j == 17 else "General"
    heat(summary, f"D2:G{1 + len(pairs)}")
    summary.freeze_panes = "D2"

    # ---- By Day: one row per session and pair (Haas's Master Summary), formulas over that session's rows
    heads = ["Date", "ETF", "Underlying"] + [f"ρ {m}m" for m in BY_DAY] + [f"β {m}m" for m in BY_DAY] \
        + [f"n {m}m" for m in BY_DAY]
    for j, h in enumerate(heads, start=1):
        header(byday, 1, j, h, bold, left=j <= 3, width=11)
    where = {}
    for s, d0 in enumerate(days):
        for k, (e, u) in enumerate(pairs):
            r = 2 + s * len(pairs) + k
            where[(d0, k)] = r
            c = byday.cell(r, 1, d0)
            c.font, c.number_format = base, "yyyy-mm-dd"
            byday.cell(r, 2, e).font = base
            byday.cell(r, 3, u).font = base
            _, _, _, a, b = pair_cols(k)
            for j, m in enumerate(BY_DAY):
                nb = len(labels(m))
                lo, hi = 2 + s * nb, 1 + (s + 1) * nb
                for off, f, fmt in ((0, "CORREL", "0.000"), (len(BY_DAY), "SLOPE", "0.000")):
                    c = byday.cell(r, 4 + off + j, f'=IFERROR({f}({rng(m, a, lo, hi)},{rng(m, b, lo, hi)}),"")')
                    c.font, c.number_format = base, fmt
                c = byday.cell(r, 4 + 2 * len(BY_DAY) + j, f"=COUNT({rng(m, a, lo, hi)})")
                c.font, c.number_format = base, "#,##0"
    n_rows = 1 + len(days) * len(pairs)
    heat(byday, f"D2:{get_column_letter(3 + len(BY_DAY))}{n_rows}")
    byday.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{n_rows}"
    byday.freeze_panes = "D2"

    # ---- 1m / 5m / 15m matrices (Haas's layout): pair x session, pointing at By Day
    for j, m in enumerate(BY_DAY):
        ws = matrices[m]
        header(ws, 1, 1, "ETF", bold, left=True, width=9)
        header(ws, 1, 2, "Underlying", bold, left=True, width=11)
        for s, d0 in enumerate(days):
            c = header(ws, 1, 3 + s, d0, bold, width=11)
            c.number_format = "yyyy-mm-dd"
        for k, (e, u) in enumerate(pairs):
            ws.cell(2 + k, 1, e).font = base
            ws.cell(2 + k, 2, u).font = base
            for s, d0 in enumerate(days):
                c = ws.cell(2 + k, 3 + s, f"='By Day'!{get_column_letter(4 + j)}{where[(d0, k)]}")
                c.font, c.number_format = base, "0.000"
        heat(ws, f"C2:{get_column_letter(2 + len(days))}{1 + len(pairs)}")
        ws.freeze_panes = "C2"

    # ---- Epps: pooled ρ and β as the sampling interval lengthens
    header(epps, 1, 1, "ρ", bold, left=True, width=13)
    header(epps, 3 + len(pairs), 1, "β", bold, left=True)
    for j, m in enumerate(MINUTES, start=2):
        for r0 in (1, 3 + len(pairs)):
            header(epps, r0, j, m, bold, width=8)
    for k, (e, u) in enumerate(pairs):
        _, _, _, a, b = pair_cols(k)
        for r0, f in ((2, "CORREL"), (4 + len(pairs), "SLOPE")):
            epps.cell(r0 + k, 1, f"{e}/{u}").font = base
            for j, m in enumerate(MINUTES, start=2):
                c = epps.cell(r0 + k, j, f"={f}({rng(m, a)},{rng(m, b)})")
                c.font, c.number_format = base, "0.000"
    heat(epps, f"B2:{get_column_letter(1 + len(MINUTES))}{1 + len(pairs)}")
    ch = LineChart()
    ch.title = "ρ vs sampling interval (minutes)"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = -1, 0
    ch.y_axis.delete = ch.x_axis.delete = False
    ch.height, ch.width = 9, 20
    ch.add_data(Reference(epps, min_col=1, max_col=1 + len(MINUTES), min_row=2, max_row=1 + len(pairs)),
                from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(epps, min_col=2, max_col=1 + len(MINUTES), min_row=1))
    for s in ch.series:
        s.marker.symbol, s.smooth = "circle", False
    epps.add_chart(ch, f"A{2 * len(pairs) + 7}")

    # ---- Data quality: how often each leg trades, and how coarse the tick is next to a minute's move
    heads = ["Symbol", "1m bars traded", "of possible", "% of minutes traded", "median volume / traded min",
             "last price", "tick (bps)", "σ per minute (bps)", "tick / σ", "zero returns"]
    for j, h in enumerate(heads, start=1):
        header(qual, 1, j, h, bold, left=j == 1, width=13)
    minutes = len(close) - len(days)                        # 1-minute intervals Yahoo returned (half days too)
    r = 2
    for k, (e, u) in enumerate(pairs):
        pe, pu, _, a, b = pair_cols(k)
        for sym, px, ret in ((e, pe, a), (u, pu, b)):
            y = close[yahoo(sym)]
            etf = sym == e                     # a cent is a coarse tick only on a $5-13 ETF
            vals = {1: sym, 2: int(y.notna().sum()), 3: len(y), 4: f"=B{r}/C{r}",
                    5: float(volume[yahoo(sym)][y.notna()].median()), 6: float(y.dropna().iloc[-1]),
                    7: f"=0.01/F{r}*10000" if etf else None, 8: f"=SQRT(SUMSQ({rng(1, ret)})/{minutes})*10000",
                    9: f"=G{r}/H{r}" if etf else None, 10: f"=COUNTIF({rng(1, ret)},0)/COUNT({rng(1, ret)})"}
            for j, v in vals.items():
                c = qual.cell(r, j, v)
                c.font = base
                c.number_format = {2: "#,##0", 3: "#,##0", 4: "0%", 5: "#,##0", 6: "#,##0.00",
                                   7: "0.0", 8: "0.0", 9: "0.00", 10: "0%"}.get(j, "General")
            r += 1

    tmp = path.with_name(f"~{path.name}")
    wb.save(tmp)
    os.replace(tmp, path)                      # a failed run never leaves a broken workbook behind


def main() -> None:
    p = argparse.ArgumentParser(description="Intraday correlation and beta, Bear 1X ETFs vs underlyings")
    p.add_argument("--sessions", type=int, default=7, help="number of sessions (Yahoo keeps ~20 of 1m bars)")
    p.add_argument("--replay", type=Path, help="rebuild from a saved download instead of Yahoo")
    p.add_argument("--partial", action="store_true", help="include today's session while it is still open")
    p.add_argument("--out", type=Path, default=BOOK)
    p.add_argument("--open", action="store_true", help="open the workbook in Excel when done")
    args = p.parse_args()

    if args.replay:
        with open(args.replay, "rb") as f:
            raw = pickle.load(f)
    else:
        print("downloading 1m and daily bars ...")
        raw = fetch(args.sessions)
    data, days, pairs = prepare(raw, args.sessions, args.partial)
    if not args.replay:
        snapshot(raw)
    print(f"  sessions {days[0]} .. {days[-1]} ({len(days)}), pairs: {', '.join(f'{e}/{u}' for e, u in pairs)}")
    build(data, days, pairs, args.out)
    print(f"wrote {args.out}")

    if args.open:
        import xlwings as xw
        xw.Book(str(args.out)).sheets["Summary"].activate()


if __name__ == "__main__":
    main()
