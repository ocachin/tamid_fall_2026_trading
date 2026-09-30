from datetime import datetime
import os
import pandas as pd
import yfinance as yf

# Openpyxl styling modules for heatmaps & visual formatting
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------
# 1. Configuration & Ticker Mapping (from uploaded image)
# ---------------------------------------------------------
PAIRS = [
    ("AVS", "AVGO"),
    ("AAPD", "AAPL"),
    ("METD", "META"),
    ("MSFD", "MSFT"),
    ("SPDN", "SPX"),
    ("AMZD", "AMZN"),
    ("PLTD", "PLTR"),
]

# Map non-standard yfinance ticker names (SPX index is ^GSPC on Yahoo Finance)
YFINANCE_TICKER_MAP = {"SPX": "^GSPC"}


def get_yf_symbol(ticker: str) -> str:
    return YFINANCE_TICKER_MAP.get(ticker, ticker)


# ---------------------------------------------------------
# 2. Main Correlation Processing Logic
# ---------------------------------------------------------
def run_correlation_analysis():
    # Define Desktop TAMID output directory for macOS
    desktop_path = os.path.expanduser("~/Desktop")
    output_dir = os.path.join(desktop_path, "TAMID")
    os.makedirs(output_dir, exist_ok=True)

    # Output file name requested
    excel_file_path = os.path.join(output_dir, "New ETF Data.xlsx")

    intervals = ["1m", "5m", "15m"]

    # Gather all unique tickers to download in batch
    all_tickers = set()
    for etf, und in PAIRS:
        all_tickers.add(get_yf_symbol(etf))
        all_tickers.add(get_yf_symbol(und))

    print("Fetching 7-day intraday market data from Yahoo Finance...")

    data_by_interval = {}

    for interval in intervals:
        print(f" -> Downloading {interval} interval data...")
        df_raw = yf.download(
            tickers=list(all_tickers),
            period="7d",
            interval=interval,
            progress=False,
        )

        # Extract Close prices
        if "Close" in df_raw.columns.levels[0]:
            df_close = df_raw["Close"]
        else:
            df_close = df_raw["Close"]

        # Convert index to US/Eastern timezone to handle Daylight Savings automatically
        if df_close.index.tz is None:
            df_close.index = df_close.index.tz_localize("UTC").tz_convert(
                "America/New_York"
            )
        else:
            df_close.index = df_close.index.tz_convert("America/New_York")

        # Strictly filter for Regular Trading Hours (Open: 09:30 AM - Close: 16:00 PM EST/EDT)
        df_market_hours = df_close.between_time("09:30", "16:00")
        data_by_interval[interval] = df_market_hours

    records = []
    sample_rows = []  # For the "Sample Raw Data" justification tab

    # Calculate daily correlations per pair and interval
    for etf_sym, und_sym in PAIRS:
        yf_etf = get_yf_symbol(etf_sym)
        yf_und = get_yf_symbol(und_sym)

        for interval in intervals:
            df_int = data_by_interval[interval]

            if yf_etf not in df_int.columns or yf_und not in df_int.columns:
                print(
                    f"Warning: Missing data for pair ({etf_sym}, {und_sym}) at {interval}"
                )
                continue

            # Merge pair prices and drop missing timestamps
            pair_df = (
                pd.DataFrame(
                    {
                        "ETF Price": df_int[yf_etf],
                        "Underlying Price": df_int[yf_und],
                    }
                )
                .dropna()
            )

            # Group by trading date
            grouped = pair_df.groupby(pair_df.index.date)
            dates = list(grouped.groups.keys())

            for trade_date, group in grouped:
                # Calculate intraday correlation
                if len(group) >= 3:
                    corr_value = group["ETF Price"].corr(
                        group["Underlying Price"]
                    )
                else:
                    corr_value = None

                records.append(
                    {
                        "Date": trade_date.strftime("%Y-%m-%d"),
                        "ETF": etf_sym,
                        "Underlying": und_sym,
                        "Interval": interval,
                        "Correlation": (
                            round(corr_value, 4)
                            if corr_value is not None
                            else None
                        ),
                    }
                )

                # Collect sample data for the justification tab (take 5m data for the most recent date)
                if (
                    interval == "5m"
                    and trade_date == dates[-1]
                    and etf_sym in ["AAPD", "AVS", "SPDN"]
                ):
                    group_sample = group.copy()
                    group_sample["ETF % Chg"] = (
                        group_sample["ETF Price"].pct_change() * 100
                    )
                    group_sample["Underlying % Chg"] = (
                        group_sample["Underlying Price"].pct_change() * 100
                    )

                    # Limit sample to 12 intraday bars (1 hour window)
                    for ts, row in group_sample.head(12).iterrows():
                        sample_rows.append(
                            {
                                "Timestamp (EST)": ts.strftime(
                                    "%Y-%m-%d %H:%M"
                                ),
                                "ETF Pair": f"{etf_sym} / {und_sym}",
                                "Interval": interval,
                                f"{etf_sym} Price": round(row["ETF Price"], 2),
                                f"{und_sym} Price": round(
                                    row["Underlying Price"], 2
                                ),
                                f"{etf_sym} % Chg": (
                                    round(row["ETF % Chg"], 3)
                                    if pd.notnull(row["ETF % Chg"])
                                    else 0.0
                                ),
                                f"{und_sym} % Chg": (
                                    round(row["Underlying % Chg"], 3)
                                    if pd.notnull(row["Underlying % Chg"])
                                    else 0.0
                                ),
                                "Daily Correlation": round(corr_value, 4),
                            }
                        )

    df_all = pd.DataFrame(records)
    df_sample = pd.DataFrame(sample_rows)

    # ---------------------------------------------------------
    # 3. Export Formatted Excel File with Heatmaps
    # ---------------------------------------------------------
    print(f"\nExporting results to Excel: {excel_file_path}")
    with pd.ExcelWriter(excel_file_path, engine="openpyxl") as writer:

        # --- Sheet 1: Master Summary ---
        pivot_master = df_all.pivot(
            index=["Date", "ETF", "Underlying"],
            columns="Interval",
            values="Correlation",
        ).reset_index()

        ordered_cols = [
            c
            for c in ["Date", "ETF", "Underlying", "1m", "5m", "15m"]
            if c in pivot_master.columns
        ]
        pivot_master = pivot_master[ordered_cols]
        pivot_master.to_excel(
            writer, sheet_name="Master Summary", index=False
        )

        # --- Sheets 2-4: Individual Interval Matrices ---
        for interval in intervals:
            df_sub = df_all[df_all["Interval"] == interval]
            if not df_sub.empty:
                matrix = df_sub.pivot(
                    index=["ETF", "Underlying"],
                    columns="Date",
                    values="Correlation",
                ).reset_index()
                sheet_title = f"{interval} Matrix"
                matrix.to_excel(writer, sheet_name=sheet_title, index=False)

        # --- Sheet 5: Sample Raw Data Justification ---
        df_sample.to_excel(
            writer, sheet_name="Sample Raw Data", index=False
        )

    # ---------------------------------------------------------
    # 4. Apply Excel Styling & Heatmaps (openpyxl)
    # ---------------------------------------------------------
    import openpyxl

    wb = openpyxl.load_workbook(excel_file_path)

    # Red-Yellow-Green 3-color scale heatmap for correlation values (-1.0 to +1.0)
    color_scale_rule = ColorScaleRule(
        start_type="num",
        start_value=-1.0,
        start_color="F8696B",  # Red (Strong negative correlation)
        mid_type="num",
        mid_value=0.0,
        mid_color="FFEB84",  # Yellow (Uncorrelated)
        end_type="num",
        end_value=1.0,
        end_color="63BE7B",  # Green (Strong positive correlation)
    )

    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        max_row = ws.max_row
        max_col = ws.max_column

        # Format header row
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center"
            )

        # Apply borders & alignment to data cells
        for row in range(2, max_row + 1):
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border
                if (
                    isinstance(cell.value, (int, float))
                    and cell.number_format == "General"
                ):
                    cell.number_format = "0.0000"

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        # Apply Heatmaps to Correlation Data Sheets
        if "Master Summary" in sheet_name:
            # Columns D, E, F are 1m, 5m, 15m correlation values
            ws.conditional_formatting.add(
                f"D2:F{max_row}", color_scale_rule
            )

        elif "Matrix" in sheet_name:
            # Correlation values span from column C to last column
            start_col_letter = openpyxl.utils.get_column_letter(3)
            end_col_letter = openpyxl.utils.get_column_letter(max_col)
            ws.conditional_formatting.add(
                f"{start_col_letter}2:{end_col_letter}{max_row}",
                color_scale_rule,
            )

    wb.save(excel_file_path)
    print(f"\nCompleted successfully!")
    print(f"File saved to: {excel_file_path}")


if __name__ == "__main__":
    run_correlation_analysis()